"""QSAR workbook: scraped data, regressions, and one live Excel chart per descriptor and level.

Run from the project root after compile:
    .venv/bin/python figures/qsar_workbook.py figures/QSAR_workbook.xlsx [preview_dir] [level ...]
(levels default to every level in config.LEVELS)

Tabs
    ReadMe       what each descriptor is and why it matters, how the fits and charts work, what
                 each statistic means, caveats and conventions
    Performance  comparison of all QSARs: best models, descriptor x level grid (R2, Q2,
                 external RMSE), every single-descriptor fit with live links to its chart tab,
                 and the two-descriptor (MLR) models with leave-one-out validation
    Data         ln kobs and every descriptor, one row per compound and level
    <LV NN name> one tab per descriptor and level: data cells, SLOPE/INTERCEPT/RSQ formulas over
                 the fitted compounds (CORR_FIT_GROUP), and a scatter chart in Jonas's style that
                 reads those cells, so the fit line updates if a value changes
Label and legend positions come from a matplotlib twin of each chart (PNGs in preview_dir if
given): each label gets the nearest Excel position plus a manual offset, with a leader line.
"""
from __future__ import annotations

import csv
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, ".")
sys.path.insert(0, "figures")
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402

import config as C  # noqa: E402
import make_excel_chart as X  # noqa: E402
from haloaro import plotting as P  # noqa: E402

SHORT = {"b3lyp_gas": "B3gas", "b3lyp_smd": "B3SMD", "m062x_gas": "M6gas", "m062x_smd": "M6SMD"}
STYLES = {"bromobenzene": "filled_gray_circle", "chlorobenzene": "open_black_circle",
          "pbde": "filled_gray_square"}
CHART_PT = (3.5 * 72, 3.0 * 72)           # chart size in points (make_excel_chart default)
# legend top-left corner (chart fractions) for each matplotlib legend location code
LEGEND_XY = {1: (0.43, 0.06), 2: (0.19, 0.06), 3: (0.19, 0.56), 4: (0.43, 0.56),
             6: (0.19, 0.32), 7: (0.43, 0.32), 8: (0.31, 0.56), 9: (0.31, 0.06)}
Y_AXIS = {"min": -8, "max": 8, "major": 2, "format": "0"}      # shared, with room for the legend
BOLD = Font(bold=True)
HEAD = Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical="top")
SHADE = PatternFill("solid", fgColor="E7E6E6")

_MARKUP = [("ArX.-", "ArX^{•−}"), ("Ar.", "Ar^{•}"), ("Ar-", "Ar^{−}"), ("X-", "X^{−}"),
           ("e-", "e^{−}"), ("H+", "H^{+}"), ("->", "→"), ("E deg", "E°"), ("dG0", "ΔG_{0}"),
           ("dG ", "ΔG "), ("C-X", "C–X"), ("Saveant", "Savéant"), ("dSCF", "ΔSCF"),
           (" RA ", " radical anion ")]

# --------------------------------------------------------------------------- #
# Documentation shown on the ReadMe tab
# --------------------------------------------------------------------------- #
# key: (units, how it is computed, why it may matter for reduction kinetics, availability)
DOCS = {
    "LUMO_eV": ("eV", "Energy of the lowest unoccupied molecular orbital of the neutral parent ArX "
                "(lowest of the alpha/beta virtual orbitals).",
                "Frontier-orbital estimate of how easily ArX accepts an electron: a lower LUMO means "
                "easier reduction. Cheap, but virtual-orbital energies depend on the functional and "
                "on diffuse functions, so it is a qualitative proxy for the electron affinity.",
                "All compounds, all levels."),
    "LUMO_RA_eV": ("eV", "LUMO of the relaxed radical anion ArX•⁻ in whatever state it optimised to "
                   "(π, bent σ, or dissociated Ar•···X⁻ complex).",
                   "How easily the radical anion takes a second electron (2-electron pathway).",
                   "All compounds; meaning changes when the radical anion is dissociated."),
    "min_dG_1e_kcal": ("kcal/mol", "ΔG of ArX + e⁻ → Ar• + X⁻ at the most favourable C–X site "
                       "(concerted dissociative one-electron reduction).",
                       "Driving force of the concerted dissociative electron transfer (Savéant). "
                       "Its standard potential is E°_DET.", "All compounds, all levels."),
    "min_dG_2e_carbanion_kcal": ("kcal/mol", "ΔG of ArX + 2e⁻ → Ar⁻ + X⁻ at the most favourable site.",
                                 "Overall driving force of two-electron reductive C–X cleavage to the "
                                 "aryl carbanion (which is protonated to ArH).",
                                 "All compounds, all levels."),
    "dG_ET_any_kcal": ("kcal/mol", "ΔG of ArX + e⁻ → ArX•⁻, using the lowest radical-anion energy "
                       "whatever its structure.",
                       "Driving force of the first electron transfer (stepwise pathway). Where the "
                       "radical anion is dissociated this already includes C–X cleavage.",
                       "All compounds, all levels."),
    "min_dG_frag_any_kcal": ("kcal/mol", "ΔG of ArX•⁻ → Ar• + X⁻ from the lowest radical-anion energy, "
                             "most favourable site.",
                             "Driving force of radical-anion fragmentation (second step of the "
                             "stepwise pathway).", "All compounds, all levels."),
    "min_dG_RA_2e_any_kcal": ("kcal/mol", "ΔG of ArX•⁻ + e⁻ → Ar⁻ + X⁻ (second electron to the radical "
                              "anion with C–X cleavage), most favourable site.",
                              "Driving force of the second reduction in a 2-electron stepwise path.",
                              "All compounds, all levels."),
    "n_X": ("-", "Number of halogen substituents.",
            "Structural baseline: most electronic descriptors increase with n_X, so a descriptor "
            "that does no better than n_X adds little physical insight.", "All compounds."),
    "VEA_eV": ("eV", "Vertical electron affinity E(ArX) − E(ArX•⁻), both at the neutral geometry "
               "(ΔSCF single points; electronic energies).",
               "Energy released on adding an electron without nuclear relaxation, i.e. the "
               "Franck–Condon step of heterogeneous electron transfer. Positive = electron is bound.",
               "All compounds, all levels (sp stage)."),
    "AEA_elec_eV": ("eV", "Adiabatic electron affinity E(ArX) − E(ArX•⁻), each at its own optimised "
                    "geometry (electronic energies).",
                    "Thermodynamic electron affinity including relaxation.",
                    "Only where the radical anion is bound (π or bent σ): few bromobenzenes at M06-2X."),
    "omega_dSCF_eV": ("eV", "Parr electrophilicity index ω = μ²/(2η) with μ = −(VIE + VEA)/2 and "
                      "η = VIE − VEA from ΔSCF vertical ionisation energy and electron affinity.",
                      "Global propensity to accept electron density; combines electron affinity and "
                      "hardness.", "All compounds, all levels (sp stage)."),
    "E_ET_V": ("V vs SHE", "Standard potential of ArX/ArX•⁻, E° = −ΔG_ET/F − E_abs(SHE).",
               "Electron-transfer potential for the stepwise pathway; compare with the −2.0 V "
               "electrode potential.", "SMD levels only, bound radical anions only."),
    "dG_ET_kcal": ("kcal/mol", "As dG_ET_any, but blank when the radical anion is dissociated.",
                   "Driving force of forming a genuine ArX•⁻ intermediate.",
                   "Bound radical anions only (often too few points to fit)."),
    "min_dG_frag_kcal": ("kcal/mol", "As min_dG_frag_any, bound radical anions only.",
                         "Fragmentation of a genuine intermediate.", "Bound radical anions only."),
    "min_dG_RA_2e_kcal": ("kcal/mol", "As min_dG_RA_2e_any, bound radical anions only.", "",
                          "Bound radical anions only."),
    "min_dG_2e_HDH_kcal": ("kcal/mol", "ΔG of ArX + H⁺ + 2e⁻ → ArH + X⁻ (hydrodehalogenation), "
                           "most favourable site.",
                           "Overall driving force of the observed net reaction (the aryl carbanion is "
                           "protonated).", "All compounds, all levels."),
    "min_dG_rad_red_kcal": ("kcal/mol", "ΔG of Ar• + e⁻ → Ar⁻ for the aryl radical formed at the most "
                            "favourable site.",
                            "Reducibility of the aryl radical: decides whether the radical is reduced "
                            "further (2e path) or reacts otherwise (H-abstraction, dimerisation).",
                            "All compounds, all levels."),
    "lambda_i_kcal": ("kcal/mol", "Nelsen four-point inner-sphere reorganisation energy of the "
                      "ArX/ArX•⁻ couple, λ_i = [E(N@A) − E(N@N)] + [E(A@N) − E(A@A)].",
                      "Inner-sphere contribution to the Marcus electron-transfer barrier (λ/4 at zero "
                      "driving force).", "π radical anions only."),
    "min_BDE_kcal": ("kcal/mol", "Weakest homolytic C–X bond enthalpy D = H(Ar•) + H(X•) − H(ArX), "
                     "with the ²P₃/₂ spin–orbit correction of X•.",
                     "Bond strength that the concerted pathway must overcome; D enters the Savéant "
                     "intrinsic barrier.", "All compounds, all levels."),
    "min_dG0_act_concerted_kcal": ("kcal/mol", "Savéant intrinsic barrier ΔG₀‡ = (D + λ₀)/4 for the "
                                   "weakest bond; λ₀ from the Marcus–Hush one-sphere model.",
                                   "Barrier of concerted dissociative electron transfer at zero "
                                   "driving force.", "SMD levels only."),
    "eff_dG_act_concerted_kcal @ -2.00 V": ("kcal/mol", "Savéant concerted barrier "
                                            "ΔG‡(E) = ΔG₀‡(1 + ΔG°(E)/4ΔG₀‡)² at E = −2.0 V, combined "
                                            "over sites as −RT ln Σ g_i exp(−ΔG‡_i/RT).",
                                            "Predicted activation free energy of the concerted pathway at "
                                            "the experimental potential; transition-state theory predicts "
                                            "a slope of −1/RT (see slope×RT).", "SMD levels only."),
    "eff_dG_act_stepwise_kcal @ -2.00 V": ("kcal/mol", "Stepwise barrier: max of the ET barrier and ET + "
                                           "fragmentation TS, degeneracy-weighted.", "",
                                           "Needs the TS stage (not run for these compounds)."),
    "eff_dG_act_combined_kcal @ -2.00 V": ("kcal/mol", "Effective barrier combining the stepwise and "
                                           "concerted rates.",
                                           "Without TS barriers this equals the concerted barrier.",
                                           "SMD levels only."),
    "eff_dG_act_frag_TS_kcal": ("kcal/mol", "C–X cleavage TS barrier of the radical anion.", "",
                                "Needs the TS stage (not run)."),
    "max_fplus_CX": ("e", "Largest condensed Fukui function f⁺ on a C–X bond (C + X), "
                     "f⁺ = q_Hirshfeld(ArX) − q_Hirshfeld(ArX•⁻ at the neutral geometry).",
                     "How much of the incoming electron lands on a C–X bond: a site-reactivity index "
                     "for reductive cleavage.", "All compounds, all levels (sp stage)."),
    "min_wiberg_CX_parent": ("-", "Smallest Wiberg C–X bond index of the neutral parent (NBO, NAO "
                             "basis).", "Weaker (lower-order) C–X bonds should break more easily.",
                             "All compounds, all levels (sp stage)."),
    "max_spin_RA_X": ("e", "Largest Hirshfeld spin density on a halogen in the relaxed radical anion.",
                      "σ* character: spin on X means the extra electron sits in the C–X σ* orbital.",
                      "Bound radical anions only (sp of the radical anion)."),
    "max_dr_CX_RA_A": ("Å", "Largest C–X lengthening in the relaxed radical anion relative to the "
                       "parent.", "Degree of C–X σ* occupation; very large values mean the radical "
                       "anion dissociated.", "All compounds, all levels."),
}

STATS = [
    ("n", "Number of compounds in the fit.",
     "Bromobenzene fits have n = 7; chlorobenzenes n = 5; halobenzenes n = 12; all n = 14.",
     "With n = 7, one unusual compound can dominate a fit."),
    ("slope, intercept", "Ordinary least squares: ln k_obs = slope·x + intercept.",
     "The slope sign shows the direction: for ΔG descriptors a negative slope means more exergonic "
     "→ faster.", "Units: ln(h⁻¹) per descriptor unit."),
    ("r, R²", "Pearson correlation and coefficient of determination R² = 1 − SSE/SST, "
     "SST = Σ(ln k − mean)².", "Fraction of the variance of ln k_obs in the fitted set explained "
     "by the line (0 to 1).", "Describes the fit, not its predictive power; R² never decreases when "
     "predictors are added."),
    ("p", "Two-sided t-test of slope = 0 with n − 2 degrees of freedom.",
     "Probability of an |r| this large if there were no linear relationship.",
     "Not corrected for multiple testing: this workbook tests ~70 descriptor × level "
     "combinations, so about 3–4 would reach p < 0.05 by chance. The Bonferroni column uses "
     "0.05/(number of fits)."),
    ("SSE, MSE, RMSE", "SSE = Σ residual²; MSE = SSE/n; RMSE = √MSE (in ln k units).",
     "Typical error of the fit. RMSE = 1 means a typical factor of e ≈ 2.7 in k_obs; RMSE = 0.5 "
     "means a factor of 1.6.", "In-sample: optimistic."),
    ("LOO RMSE, Q²", "Leave-one-out cross-validation: refit without each compound in turn and "
     "predict it; PRESS = Σ(prediction error)²; LOO RMSE = √(PRESS/n); Q² = 1 − PRESS/SST.",
     "Predictive ability within the fitted set. Q² close to R² means a robust fit; Q² much lower "
     "than R² means one compound is driving the fit; Q² < 0 means the model predicts worse than "
     "the mean.", "Preferred ranking statistic here (Performance tab ranks by Q²)."),
    ("ext RMSE", "RMSE of ln k_obs predicted by the bromobenzene fit for the compounds not used in "
     "the fit (chlorobenzenes and PBDEs).", "External validation: does the relationship transfer "
     "to other halogens and to the diphenyl ethers?", "A systematic offset (e.g. all PBDEs slower) "
     "inflates it even when the ranking is right."),
    ("R² halobenzenes / all", "R² of separate fits on the 12 bromo- plus chlorobenzenes, or all 14 "
     "compounds.", "Whether one line describes both halogens (and the PBDEs).",
     "These fits use more data but mix chemically different families."),
    ("adj R² (MLR)", "1 − (1 − R²)(n − 1)/(n − p − 1), p = 2 predictors.",
     "R² penalised for the extra parameter; compare MLR and single-descriptor models with adj R² "
     "and Q², not R².", "With n = 7 and 3 parameters only 4 degrees of freedom remain."),
    ("slope×RT", "−slope·RT for barrier descriptors (RT = 0.593 kcal/mol at 298 K).",
     "Transition-state theory predicts ln k = const − ΔG‡/RT, i.e. slope×RT = 1.",
     "Values far from 1 mean the barrier descriptor scales but does not match TST quantitatively."),
]

CAVEATS = [
    "Small data set: 7 bromobenzenes define each fit. Treat R² > 0.9 with n = 7 as suggestive, and "
    "rely on Q² and the external RMSE for predictive claims.",
    "Multiple testing: many descriptors and four levels were screened. The best-looking fit is "
    "partly selected by chance; prefer descriptors that do well at several levels and in the "
    "external test.",
    "Collinearity: most descriptors correlate with the number of halogens (n_X). A descriptor that "
    "does not beat n_X adds little.",
    "Radical-anion state: at these levels most radical anions are bent σ-type or dissociated "
    "(Ar•···X⁻ complexes). *_any descriptors use the lowest radical-anion energy whatever its state; "
    "the 'bound radical anion only' descriptors are blank where it dissociated, so they have too "
    "few points to fit.",
    "Stepwise barriers need the TS stage, which was not run: C–X scans of the BDE-99 radical anion "
    "showed electronic-state crossings rather than saddle points.",
    "B3LYP/SMD BDE-99 radical anion did not finish (optimisation step limit), so the *_any "
    "descriptors of BDE-99 are missing at B3LYP/SMD.",
    "Chlorobenzene rate constants are unpublished and have no standard errors yet.",
    "Gas-phase vs SMD: the cathode is activated carbon, so neither gas phase nor bulk water is the "
    "true environment; compare both.",
]


# --------------------------------------------------------------------------- #
def markup(text):
    """Config label -> make_excel_chart markup (^{} superscript, _{} subscript)."""
    text = f" {text} "
    for a, b in _MARKUP:
        text = text.replace(a, b)
    return text.strip()


def x_title(desc, label):
    at = re.search(r"@ ([+-]?\d+\.\d+) V", desc)
    t = label + (f" at {float(at.group(1)):.1f} V".replace("-", "−") if at else "")
    u = P.unit(desc).replace("$^{3}$", "^{3}")
    return markup(t) + (f" ({u})" if u else "")


def excel_position(dx, dy):
    """Excel label position closest to a matplotlib offset (points), and that position's own
    offset, so the remainder can be applied as a manual layout shift."""
    if abs(dx) > 2:
        return ("right", (6, 0)) if dx > 0 else ("left", (-6, 0))
    return ("above", (0, 8)) if dy > 0 else ("below", (0, -8))


def sheet_name(lv, rank, desc, used):
    slug = desc.split(" @")[0]
    for a in ("min_", "max_", "eff_", "_kcal", "_eV", "_A", "dG_"):
        slug = slug.replace(a, "" if a != "dG_" else "dG ")
    name = re.sub(r"[\[\]:*?/\\]", "", f"{SHORT[lv]} {rank:02d} {slug}")[:31]
    while name in used:
        name = name[:29] + f"{len(used) % 100:02d}"
    used.add(name)
    return name


def twin_layout(data, lv, desc, fit, label, preview):
    """Draw the chart in matplotlib to choose label positions and the legend corner."""
    pts = P._points(data, lv, desc)
    fig, ax = plt.subplots(figsize=(3.5, 3.0))
    yt = [float(v) for v in range(Y_AXIS["min"], Y_AXIS["max"] + 1, Y_AXIS["major"])]
    drawn, line = P._panel(ax, fit, pts, desc, label, (yt, 0), ms=7)
    fig.tight_layout()
    _, leg = P._legend(fig, ax, drawn, pts, line, 9)
    placed = P._place_labels(fig, ax, pts, line, avoid=[leg])
    if preview:
        fig.savefig(preview, dpi=150, facecolor="white")
    plt.close(fig)
    return pts, placed, leg._loc


def chart_spec(desc, label, pts, placed, loc, caption):
    series = []
    for grp in P.ROLES:
        sel = [p for p in pts if p["group"] == grp]
        if not sel:
            continue
        pos, offs = [], []
        for p in sel:
            dx, dy = placed.get(p["label"], (6, 0))
            where, (bx, by) = excel_position(dx, dy)
            pos.append(where)
            ox, oy = (dx - bx) / CHART_PT[0], -(dy - by) / CHART_PT[1]
            offs.append([round(ox, 4), round(oy, 4)] if abs(dx - bx) > 2 or abs(dy - by) > 2 else None)
        s = {"name": P.ROLES[grp]["label"], "x_header": desc, "x": [p["x"] for p in sel],
             "y": [round(p["y"], 4) for p in sel], "labels": [p["label"] for p in sel],
             "label_positions": pos, "label_offsets": offs, "style": STYLES[grp]}
        if all(p["yerr"] is not None for p in sel):
            s["y_err"] = [round(p["yerr"], 4) for p in sel]
        if grp == C.CORR_FIT_GROUP:
            s["trendline"] = "linear"
        series.append(s)
    ticks, dec = P.nice_ticks([p["x"] for p in pts])
    lx, ly = LEGEND_XY.get(loc, (0.43, 0.06))
    return {"x_title": x_title(desc, label), "y_title": "ln k_{obs} (h^{−1})",
            "x_axis": {"min": ticks[0], "max": ticks[-1], "major": round(ticks[1] - ticks[0], 10),
                       "format": "0" if dec == 0 else "0." + "0" * dec},
            "y_axis": Y_AXIS, "series": series, "legend": {"x": lx, "y": ly}, "caption": caption}


def mlr_loo(x1, x2, y):
    """R2, adj R2, RMSE, LOO RMSE and Q2 of y = b0 + b1 x1 + b2 x2."""
    X_ = np.column_stack([np.ones(len(y)), x1, x2])
    y = np.asarray(y)
    b = np.linalg.lstsq(X_, y, rcond=None)[0]
    sse = float(((y - X_ @ b) ** 2).sum())
    sst = float(((y - y.mean()) ** 2).sum())
    press = 0.0
    for i in range(len(y)):
        m = np.arange(len(y)) != i
        bi = np.linalg.lstsq(X_[m], y[m], rcond=None)[0]
        press += float((y[i] - X_[i] @ bi) ** 2)
    n = len(y)
    r2 = 1 - sse / sst
    return {"n": n, "R2": r2, "adj_R2": 1 - (1 - r2) * (n - 1) / (n - 3), "RMSE": math.sqrt(sse / n),
            "LOO_RMSE": math.sqrt(press / n), "Q2": 1 - press / sst, "b": b}


def num(v):
    return float(v) if v not in (None, "") else None


def write_rows(ws, rows, start_row, bold_first=True, widths=None):
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=start_row + i, column=j + 1, value=v)
            if i == 0 and bold_first:
                c.font = BOLD
                c.fill = SHADE
            c.alignment = WRAP
    return start_row + len(rows)


def fit_heights(ws, first=1, last=None, line_pt=13.5):
    """Row heights from the wrapped text length and column widths (Excel does not re-fit rows
    written by openpyxl). Merged cells count with the width of the whole merged range."""
    merged = {}
    for rng in ws.merged_cells.ranges:
        w = sum((ws.column_dimensions[get_column_letter(c)].width or 10) for c in range(rng.min_col, rng.max_col + 1))
        merged[(rng.min_row, rng.min_col)] = w
    for row in ws.iter_rows(min_row=first, max_row=last or ws.max_row):
        lines = 1
        for c in row:
            if isinstance(c.value, str) and c.alignment is not None and c.alignment.wrap_text:
                w = merged.get((c.row, c.column), ws.column_dimensions[c.column_letter].width or 10)
                lines = max(lines, sum(math.ceil(max(1, len(part)) / max(1.0, w * 1.1)) for part in c.value.split("\n")))
        if lines > 1:
            ws.row_dimensions[row[0].row].height = lines * line_pt


# --------------------------------------------------------------------------- #
def main(out, preview_dir=None, levels=None):
    levels = levels or list(C.LEVELS)
    data = list(csv.DictReader(open("results/correlation_data.csv")))
    qsar = list(csv.DictReader(open("results/qsar_summary.csv")))
    labels = {d: lab for d, lab, _ in C.QSAR_DESCRIPTORS}
    sets = {d: s for d, _, s in C.QSAR_DESCRIPTORS}
    fits = {(r["level"], r["fit_set"], r["descriptors"]): r for r in qsar}
    if preview_dir:
        Path(preview_dir).mkdir(parents=True, exist_ok=True)
    P.use_style()
    st = dict(X.STYLE)

    wb = Workbook()
    readme = wb.active
    readme.title = "ReadMe"
    perf = wb.create_sheet("Performance")
    data_ws = wb.create_sheet("Data")

    # ---------------- chart tabs ----------------
    jobs, rows, used, first = {}, [], set(), True
    for lv in levels:
        best = [r for r in qsar if r["level"] == lv and r["fit_set"] == C.CORR_FIT_GROUP
                and r["descriptor_set"] != "MLR" and r.get("R2") not in (None, "")]
        best.sort(key=lambda r: -float(r["R2"]))
        for rank, fit in enumerate(best, 1):
            desc = fit["descriptors"]
            name = sheet_name(lv, rank, desc, used)
            preview = f"{preview_dir}/{name.replace(' ', '_')}.png" if preview_dir else None
            pts, placed, loc = twin_layout(data, lv, desc, fit, labels[desc], preview)
            cap = (P._caption("", lv, [fit], first, True, {p["group"] for p in pts}).strip()
                   + " Fit line, slope, intercept and R² are Excel formulas (SLOPE, INTERCEPT, RSQ) "
                     "over the bromobenzene cells beside the chart.")
            first = False
            ws = wb.create_sheet(name)
            jobs[name] = X.add_chart_sheet(ws, chart_spec(desc, labels[desc], pts, placed, loc, cap), st)
            cells = {}
            for row in ws.iter_rows(min_row=5, max_row=7):
                for c in row:
                    if c.value in ("slope", "intercept", "R2"):
                        cells[c.value] = f"'{name}'!{get_column_letter(c.column + 1)}{c.row}"
            rows.append((lv, rank, name, desc, fit, cells))

    n_fits = len(rows)
    bonf = 0.05 / max(1, n_fits)

    # ---------------- Performance ----------------
    perf["A1"] = "QSAR performance: ln k_obs vs computed descriptors (single-descriptor fits on the bromobenzenes, n = 7)"
    perf["A1"].font = HEAD
    perf["A2"] = ("Ranked by leave-one-out Q² (predictive ability). R² describes the fit; Q² tests it "
                  "by predicting each bromobenzene from the other six; ext RMSE applies the bromobenzene "
                  "line to the chlorobenzenes and PBDEs, which were not used in the fit. RMSE values are in "
                  f"ln k units (1 = factor 2.7 in k_obs). Bonferroni p threshold for {n_fits} fits: {bonf:.1e}. "
                  "Definitions on the ReadMe tab.")
    perf["A2"].alignment = WRAP
    perf.merge_cells("A2:N2")
    perf.row_dimensions[2].height = 45

    single = []
    for lv, rank, name, desc, fit, cells in rows:
        oth = {fs: fits.get((lv, fs, desc), {}) for fs in ("chlorobenzene", "halobenzene", "all")}
        single.append({"level": lv, "name": name, "desc": desc, "label": P.chem(labels[desc], plain=True),
                       "set": sets[desc], "n": int(fit["n"]), "R2": num(fit["R2"]), "Q2": num(fit.get("Q2")),
                       "RMSE": num(fit["RMSE"]), "LOO": num(fit.get("LOO_RMSE")), "p": num(fit["p"]),
                       "slope": num(fit["slope"]), "intercept": num(fit["intercept"]),
                       "ext": num(fit.get("pred_RMSE_other_groups")), "slopeRT": num(fit.get("slope_x_RT")),
                       "R2_cl": num(oth["chlorobenzene"].get("R2")), "R2_hal": num(oth["halobenzene"].get("R2")),
                       "R2_all": num(oth["all"].get("R2")), "cells": cells})
    single.sort(key=lambda r: -(r["Q2"] if r["Q2"] is not None else -9))

    r0 = 4
    perf.cell(row=r0, column=1, value="1. Top 10 models (all levels)").font = HEAD
    hdr = ["rank", "level", "descriptor", "R²", "Q²", "RMSE", "LOO RMSE", "ext RMSE", "p",
           "p < Bonferroni", "R² halobenzenes (n=12)", "chart tab"]
    tbl = [hdr] + [[i + 1, SHORT[r["level"]], r["label"], r["R2"], r["Q2"], r["RMSE"], r["LOO"], r["ext"],
                    r["p"], "yes" if r["p"] is not None and r["p"] < bonf else "no", r["R2_hal"], r["name"]]
                   for i, r in enumerate(single[:10])]
    end = write_rows(perf, tbl, r0 + 1)
    for i in range(10):
        c = perf.cell(row=r0 + 2 + i, column=12)
        c.hyperlink, c.font = f"#'{c.value}'!A1", Font(color="0563C1", underline="single")

    # descriptor x level grid
    r0 = end + 2
    perf.cell(row=r0, column=1, value="2. Descriptor × level comparison (bromobenzene fits)").font = HEAD
    perf.cell(row=r0 + 1, column=1, value="Bold = best value in that column. Blank = too few compounds "
              "with data at that level.").alignment = WRAP
    descs = [d for d, _, _ in C.QSAR_DESCRIPTORS if any(r["desc"] == d for r in single)]
    hdr1 = ["descriptor", "set"] + [f"{SHORT[lv]} {s}" for s in ("R²", "Q²", "ext RMSE") for lv in levels] \
        + [f"{SHORT[lv]} R² halo" for lv in levels]
    by = {(r["level"], r["desc"]): r for r in single}
    grid = [hdr1]
    for d in descs:
        row = [P.chem(labels[d], plain=True), sets[d]]
        for key in ("R2", "Q2", "ext"):
            row += [by.get((lv, d), {}).get(key) for lv in levels]
        row += [by.get((lv, d), {}).get("R2_hal") for lv in levels]
        grid.append(row)
    top = r0 + 2
    end = write_rows(perf, grid, top)
    for j in range(2, len(hdr1)):
        vals = [(perf.cell(row=top + 1 + i, column=j + 1).value, i) for i in range(len(descs))]
        vals = [(v, i) for v, i in vals if isinstance(v, (int, float))]
        if vals:
            pick = min(vals) if "ext" in hdr1[j] else max(vals)
            perf.cell(row=top + 1 + pick[1], column=j + 1).font = BOLD

    # all single fits
    r0 = end + 2
    perf.cell(row=r0, column=1, value="3. All single-descriptor fits (sorted by Q²)").font = HEAD
    hdr = ["level", "descriptor", "key", "set", "n", "R² (live)", "slope (live)", "R²", "Q²", "slope",
           "intercept", "p", "RMSE", "LOO RMSE", "ext RMSE", "slope×RT", "R² chlorobenzenes (n=5)",
           "R² halobenzenes (n=12)", "R² all (n=14)", "chart tab"]
    tbl = [hdr]
    for r in single:
        tbl.append([SHORT[r["level"]], r["label"], r["desc"], r["set"], r["n"],
                    f"={r['cells']['R2']}" if "R2" in r["cells"] else None,
                    f"={r['cells']['slope']}" if "slope" in r["cells"] else None,
                    r["R2"], r["Q2"], r["slope"], r["intercept"], r["p"], r["RMSE"], r["LOO"], r["ext"],
                    r["slopeRT"], r["R2_cl"], r["R2_hal"], r["R2_all"], r["name"]])
    start = r0 + 1
    end = write_rows(perf, tbl, start)
    for i in range(len(single)):
        c = perf.cell(row=start + 1 + i, column=20)
        c.hyperlink, c.font = f"#'{c.value}'!A1", Font(color="0563C1", underline="single")

    # MLR
    r0 = end + 2
    perf.cell(row=r0, column=1, value="4. Two-descriptor models (multiple linear regression)").font = HEAD
    perf.cell(row=r0 + 1, column=1, value="ln k = b0 + b1·x1 + b2·x2. Compare with single-descriptor "
              "models through adj R² and Q² (R² always rises with an extra predictor).").alignment = WRAP
    hdr = ["level", "fit set", "x1", "x2", "n", "R²", "adj R²", "Q²", "RMSE", "LOO RMSE", "b0", "b1", "b2"]
    tbl = [hdr]
    D = {(r["level"], r["species"]): r for r in data}
    for lv in levels:
        for a, b in C.CORR_PAIRS:
            for fs in C.CORR_FIT_SETS:
                from haloaro.correlations import in_set
                pts = [(num(r.get(a)), num(r.get(b)), num(r["ln_kobs"])) for r in data
                       if r["level"] == lv and in_set(r, fs)]
                pts = [p for p in pts if None not in p]
                if len(pts) < 5:
                    continue
                m = mlr_loo(*zip(*pts))
                tbl.append([SHORT[lv], fs, P.chem(labels.get(a, a), plain=True), P.chem(labels.get(b, b), plain=True),
                            m["n"], m["R2"], m["adj_R2"], m["Q2"], m["RMSE"], m["LOO_RMSE"], *map(float, m["b"])])
    end = write_rows(perf, tbl, r0 + 2)
    for col, w in zip(range(1, 21), (8, 40, 30, 10, 8, 11, 11, 9, 9, 10, 10, 10, 9, 10, 10, 10, 13, 13, 11, 24)):
        perf.column_dimensions[get_column_letter(col)].width = w
    for row in perf.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, float) or (isinstance(c.value, str) and c.value.startswith("='")):
                c.number_format = "0.000" if not (isinstance(c.value, float) and abs(c.value) < 1e-3 and c.value) else "0.0E+00"
    perf.freeze_panes = "A4"

    # ---------------- Data ----------------
    dcols = [d for d, _, _ in C.QSAR_DESCRIPTORS if any(r.get(d) not in (None, "") for r in data
                                                         if r["level"] in levels)]
    write_rows(data_ws, [["level", "compound", "label", "species", "group", "kobs (1/h)", "SE (1/h)", "ln kobs"]
                         + dcols], 1)
    for r in data:
        if r["level"] in levels:
            data_ws.append([SHORT[r["level"]], r["compound"], r["label"], r["species"], r["group"],
                            num(r["kobs_per_h"]), num(r.get("se_per_h")), num(r["ln_kobs"])]
                           + [num(r.get(d)) for d in dcols])
    data_ws.freeze_panes = "C2"
    for i in range(1, data_ws.max_column + 1):
        data_ws.column_dimensions[get_column_letter(i)].width = 14

    # ---------------- ReadMe ----------------
    lvtext = "; ".join(f"{SHORT[lv]} = {P.level_text(lv)}" for lv in levels)
    blocks = [
        ("Haloaromatics QSAR workbook", None),
        ("Purpose", "Correlate measured electrochemical dehalogenation rate constants with computed "
                    "descriptors, to identify which parameters predict reduction kinetics and whether a "
                    "relationship fitted on the bromobenzenes transfers to the chlorobenzenes and PBDEs."),
        ("Rate constants", f"Pseudo-first-order k_obs (h⁻¹) measured at {P.FIRST_E}. Bromobenzenes and PBDEs: "
                           "LaPier et al., Environ. Sci. Technol. 2026, 60, 1346, Table 1 (1,3-dibromobenzene "
                           "corrected from per day to 0.0329 h⁻¹). Chlorobenzenes: unpublished, no standard "
                           "errors yet. File: data/experimental_kobs.csv."),
        ("Computed data", f"Gaussian 16, 6-311++G(d). Levels: {lvtext}. Compounds: 7 bromobenzenes, 5 "
                          "chlorobenzenes, BDE-47 and BDE-99. Built from results/qsar_summary.csv and "
                          "correlation_data.csv (compile) by figures/qsar_workbook.py."),
        ("Tabs", "Performance: comparison of all models. Data: every value used. One chart tab per "
                 "descriptor and level, named <level> <rank by R²> <descriptor>: B3 = B3LYP, M6 = M06-2X, "
                 "gas or SMD."),
        ("Charts and fits", "Each chart tab holds the plotted values in cells. The dotted line is the "
                            "least-squares fit to the bromobenzenes computed by Excel (SLOPE, INTERCEPT, "
                            "RSQ in the cells beside the chart), so editing a value updates the line and "
                            "statistics. Filled gray circles: published bromobenzene data (fitted). Open "
                            "circles: new chlorobenzene data (predicted). Gray squares: published PBDE "
                            "data (predicted). Error bars: ±1 standard error of k_obs propagated to ln k_obs "
                            "(SE/k_obs). Site descriptors use the most favourable C–X site of each compound "
                            "(min_ / max_) or a degeneracy-weighted effective value (eff_)."),
    ]
    r = 1
    readme.cell(row=r, column=1, value=blocks[0][0]).font = Font(bold=True, size=14)
    r += 2
    for title, text in blocks[1:]:
        readme.cell(row=r, column=1, value=title).font = BOLD
        c = readme.cell(row=r, column=2, value=text)
        c.alignment = WRAP
        readme.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        r += 1
    r += 1
    readme.cell(row=r, column=1, value="Calculated parameters").font = HEAD
    r += 1
    prm = [["descriptor", "key", "units", "definition / how computed", "significance for reduction kinetics",
            "availability", "set"]]
    for d, lab, s in C.QSAR_DESCRIPTORS:
        u, how, why, avail = DOCS.get(d, ("", "", "", ""))
        prm.append([P.chem(lab, plain=True), d, u, how, why, avail, s])
    r = write_rows(readme, prm, r) + 1
    readme.cell(row=r, column=1, value="Statistics").font = HEAD
    r += 1
    r = write_rows(readme, [["statistic", "definition", "how to read it", "caveats"]] + [list(s) for s in STATS], r) + 1
    readme.cell(row=r, column=1, value="Caveats").font = HEAD
    r += 1
    for t in CAVEATS:
        c = readme.cell(row=r, column=2, value=t)
        c.alignment = WRAP
        readme.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        readme.cell(row=r, column=1, value="•")
        r += 1
    r += 1
    readme.cell(row=r, column=1, value="Conventions").font = HEAD
    r += 1
    conv = (f"T = {C.TEMPERATURE} K; G(e⁻) = {C.G_ELECTRON_KCAL} kcal/mol (Fermi–Dirac); E_abs(SHE) = "
            f"{C.E_ABS_SHE_V} V; SMD solutes +{C.STD_STATE_CORR_KCAL} kcal/mol (1 atm → 1 M); "
            f"potential-dependent barriers at {C.EXP_POTENTIAL_V:+.1f} V vs SHE; ln k_obs with k_obs in h⁻¹.")
    c = readme.cell(row=r, column=2, value=conv)
    c.alignment = WRAP
    readme.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    for col, w in zip("ABCDEFG", (30, 28, 10, 55, 55, 34, 9)):
        readme.column_dimensions[col].width = w

    fit_heights(readme)
    fit_heights(perf, 4)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    X.patch_many(out, st, jobs)
    print(f"wrote {out}: ReadMe, Performance, Data and {n_fits} chart tabs")


if __name__ == "__main__":
    a = sys.argv[1:]
    main(a[0] if a else "figures/QSAR_workbook.xlsx", a[1] if len(a) > 1 else None, a[2:] or None)
