"""M06-2X QSAR workbook: scraped data, regressions, and one live Excel chart per descriptor.

Run from the project root after compile:
    .venv/bin/python figures/m062x_qsar_workbook.py figures/M062X_QSAR.xlsx [preview_dir]

Tabs
    Summary   every single-descriptor fit at M06-2X gas and SMD (fit on CORR_FIT_GROUP),
              ranked by R2, with live links to the R2/slope cells of each chart tab, the other
              fit sets, and the two-descriptor (MLR) models
    Data      ln kobs and every descriptor, one row per compound and level
    <level NN name>  one tab per descriptor: data cells, SLOPE/INTERCEPT/RSQ formulas over the
              fitted compounds, and a scatter chart in Jonas's style that reads those cells.
              The fit line is drawn from the formula cells, so it updates if a value changes.
Label and legend positions are taken from a matplotlib twin of each chart (written to
preview_dir as PNGs if given): each label gets the nearest Excel position (right/left/above/
below) plus a manual offset, and Excel draws a leader line when a label sits away from its point.
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, ".")
sys.path.insert(0, "figures")
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402

import config as C  # noqa: E402
import make_excel_chart as X  # noqa: E402
from haloaro import plotting as P  # noqa: E402

LEVELS = ["m062x_gas", "m062x_smd"]
SHORT = {"m062x_gas": "gas", "m062x_smd": "SMD"}
STYLES = {"bromobenzene": "filled_gray_circle", "chlorobenzene": "open_black_circle",
          "pbde": "filled_gray_square"}
CHART_PT = (3.5 * 72, 3.0 * 72)           # chart size in points (make_excel_chart default)


def excel_position(dx, dy):
    """Excel label position closest to a matplotlib offset (points), and that position's
    own offset, so the remainder can be applied as a manual layout shift."""
    if abs(dx) > 2:
        return ("right", (6, 0)) if dx > 0 else ("left", (-6, 0))
    return ("above", (0, 8)) if dy > 0 else ("below", (0, -8))
# legend top-left corner (chart fractions) for each matplotlib legend location code
LEGEND_XY = {1: (0.43, 0.06), 2: (0.19, 0.06), 3: (0.19, 0.56), 4: (0.43, 0.56),
             6: (0.19, 0.32), 7: (0.43, 0.32), 8: (0.31, 0.56), 9: (0.31, 0.06)}
Y_AXIS = {"min": -8, "max": 8, "major": 2, "format": "0"}      # shared, with room for the legend

_MARKUP = [("ArX.-", "ArX^{•−}"), ("Ar.", "Ar^{•}"), ("Ar-", "Ar^{−}"), ("X-", "X^{−}"),
           ("e-", "e^{−}"), ("H+", "H^{+}"), ("->", "→"), ("E deg", "E°"), ("dG0", "ΔG_{0}"),
           ("dG ", "ΔG "), ("C-X", "C–X"), ("Saveant", "Savéant"), ("dSCF", "ΔSCF"),
           (" RA ", " radical anion ")]


def markup(text):
    """Config label -> make_excel_chart markup (^{} superscript, _{} subscript)."""
    text = f" {text} "
    for a, b in _MARKUP:
        text = text.replace(a, b)
    return text.strip()


def x_title(desc, label):
    at = re.search(r"@ ([+-]?\d+\.\d+) V", desc)
    t = label + (f" at {float(at.group(1)):.1f} V" if at else "")
    u = P.unit(desc).replace("$^{3}$", "^{3}")
    return markup(t) + (f" ({u})" if u else "")


def sheet_name(lv, rank, desc, used):
    slug = desc.split(" @")[0]
    for a in ("min_", "max_", "eff_", "_kcal", "_eV", "_A", "dG_"):
        slug = slug.replace(a, "" if a != "dG_" else "dG ")
    name = f"{SHORT[lv]} {rank:02d} {slug}"
    name = re.sub(r"[\[\]:*?/\\]", "", name)[:31]
    while name in used:
        name = name[:29] + f"{len(used) % 100:02d}"
    used.add(name)
    return name


def twin_layout(data, lv, desc, fit, label, preview):
    """Draw the chart in matplotlib to choose label directions and the legend corner."""
    pts = P._points(data, lv, desc)
    fig, ax = plt.subplots(figsize=(3.5, 3.0))
    yt = [float(v) for v in range(Y_AXIS["min"], Y_AXIS["max"] + 1, Y_AXIS["major"])]
    drawn, line = P._panel(ax, fit, pts, desc, label, (yt, 0), ms=7)
    fig.tight_layout()
    _, leg = P._legend(fig, ax, drawn, pts, line, 9)
    placed = P._place_labels(fig, ax, pts, line, avoid=[leg])
    if preview:
        fig.savefig(preview, dpi=150, facecolor="white")
    plt.close(fig)
    return pts, placed, leg._loc


def chart_spec(lv, desc, fit, label, pts, placed, loc, caption):
    series = []
    for grp in P.ROLES:
        sel = [p for p in pts if p["group"] == grp]
        if not sel:
            continue
        pos, offs = [], []
        for p in sel:
            dx, dy = placed.get(p["label"], (6, 0))
            where, (bx, by) = excel_position(dx, dy)
            pos.append(where)
            # Excel data-label layout: offset from the default position, in chart fractions
            ox, oy = (dx - bx) / CHART_PT[0], -(dy - by) / CHART_PT[1]
            offs.append([round(ox, 4), round(oy, 4)] if abs(dx - bx) > 2 or abs(dy - by) > 2 else None)
        s = {"name": P.ROLES[grp]["label"], "x_header": desc, "x": [p["x"] for p in sel],
             "y": [round(p["y"], 4) for p in sel], "labels": [p["label"] for p in sel],
             "label_positions": pos, "label_offsets": offs, "style": STYLES[grp]}
        if all(p["yerr"] is not None for p in sel):
            s["y_err"] = [round(p["yerr"], 4) for p in sel]
        if grp == C.CORR_FIT_GROUP:
            s["trendline"] = "linear"
        series.append(s)
    ticks, dec = P.nice_ticks([p["x"] for p in pts])
    lx, ly = LEGEND_XY.get(loc, (0.43, 0.06))
    return {"x_title": x_title(desc, label), "y_title": "ln k_{obs} (h^{−1})",
            "x_axis": {"min": ticks[0], "max": ticks[-1], "major": round(ticks[1] - ticks[0], 10),
                       "format": "0" if dec == 0 else "0." + "0" * dec},
            "y_axis": Y_AXIS, "series": series, "legend": {"x": lx, "y": ly}, "caption": caption}


def bold_row(ws, row):
    for c in ws[row]:
        c.font = Font(bold=True)


def main(out, preview_dir=None):
    data = list(csv.DictReader(open("results/correlation_data.csv")))
    qsar = list(csv.DictReader(open("results/qsar_summary.csv")))
    labels = {d: lab for d, lab, _ in C.QSAR_DESCRIPTORS}
    fits = {(r["level"], r["fit_set"], r["descriptors"]): r for r in qsar}
    if preview_dir:
        Path(preview_dir).mkdir(parents=True, exist_ok=True)
    P.use_style()
    st = dict(X.STYLE)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    data_ws = wb.create_sheet("Data")
    jobs, rows, used, first = {}, [], set(), True
    for lv in LEVELS:
        best = [r for r in qsar if r["level"] == lv and r["fit_set"] == C.CORR_FIT_GROUP
                and r["descriptor_set"] != "MLR" and r.get("R2") not in (None, "")]
        best.sort(key=lambda r: -float(r["R2"]))
        for rank, fit in enumerate(best, 1):
            desc = fit["descriptors"]
            name = sheet_name(lv, rank, desc, used)
            preview = f"{preview_dir}/{name.replace(' ', '_')}.png" if preview_dir else None
            pts, placed, loc = twin_layout(data, lv, desc, fit, labels[desc], preview)
            present = {p["group"] for p in pts}
            cap = (P._caption("", lv, [fit], first, True, present).strip()
                   + " Fit line, slope, intercept and R² are Excel formulas (SLOPE, INTERCEPT, RSQ) "
                     "over the bromobenzene cells beside the chart.")
            first = False
            spec = chart_spec(lv, desc, fit, labels[desc], pts, placed, loc, cap)
            ws = wb.create_sheet(name)
            jobs[name] = X.add_chart_sheet(ws, spec, st)
            # find the formula cells written by add_chart_sheet (label column, rows 5-7)
            cells = {}
            for row in ws.iter_rows(min_row=5, max_row=7):
                for c in row:
                    if c.value in ("slope", "intercept", "R2"):
                        cells[c.value] = f"'{name}'!{get_column_letter(c.column + 1)}{c.row}"
            rows.append((lv, rank, name, desc, fit, cells))

    # ---------------- Summary ----------------
    summary["A1"] = "M06-2X/6-311++G(d) QSAR: ln k_obs vs computed descriptors"
    summary["A1"].font = Font(bold=True, size=12)
    summary["A2"] = (f"k_obs measured at {P.FIRST_E}. Single-descriptor least-squares fits on the "
                     f"{C.CORR_FIT_GROUP}s (n = 7); the other compounds are predicted. R² and slope "
                     "(live) link to the formula cells of each chart tab. Descriptors: most favourable "
                     "site; *_any uses the lowest radical-anion energy whatever its state.")
    summary["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A2:P2")
    summary.row_dimensions[2].height = 45
    hdr = ["level", "rank", "chart tab", "descriptor", "key", "n", "R² (live)", "slope (live)",
           "R²", "slope", "intercept", "p", "RMSE", "LOO Q²", "pred RMSE (other groups)",
           "R² chlorobenzenes", "R² halobenzenes", "R² all", "n all"]
    summary.append([])
    summary.append(hdr)
    bold_row(summary, 4)
    for lv, rank, name, desc, fit, cells in rows:
        other = {fs: fits.get((lv, fs, desc), {}) for fs in ("chlorobenzene", "halobenzene", "all")}
        num = lambda v: float(v) if v not in (None, "") else None  # noqa: E731
        summary.append([SHORT[lv], rank, name, P.chem(labels[desc], plain=True), desc, int(fit["n"]),
                        f"={cells['R2']}" if "R2" in cells else None,
                        f"={cells['slope']}" if "slope" in cells else None,
                        num(fit["R2"]), num(fit["slope"]), num(fit["intercept"]), num(fit["p"]),
                        num(fit["RMSE"]), num(fit.get("Q2")), num(fit.get("pred_RMSE_other_groups")),
                        num(other["chlorobenzene"].get("R2")), num(other["halobenzene"].get("R2")),
                        num(other["all"].get("R2")), num(other["all"].get("n"))])
        c = summary.cell(row=summary.max_row, column=3)
        c.hyperlink = f"#'{name}'!A1"
        c.font = Font(color="0563C1", underline="single")
    summary.append([])
    summary.append(["Two-descriptor models (MLR, static values)"])
    bold_row(summary, summary.max_row)
    summary.append(["level", "fit set", "model", "n", "R²", "adj R²", "RMSE", "b0", "b1", "b2"])
    bold_row(summary, summary.max_row)
    for r in qsar:
        if r["level"] in LEVELS and r["descriptor_set"] == "MLR" and r.get("R2"):
            summary.append([SHORT[r["level"]], r["fit_set"], P.chem(r["model"], plain=True), int(r["n"]),
                            *[float(r[k]) if r.get(k) not in (None, "") else None
                              for k in ("R2", "adj_R2", "RMSE", "b0", "b1", "b2")]])
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", (7, 6, 24, 44, 34, 5, 10, 11, 8, 9, 10, 9, 8, 8, 12, 12, 12, 8, 6)):
        summary.column_dimensions[col].width = w
    for row in summary.iter_rows(min_row=5):
        for c in row[6:18]:
            c.number_format = "0.000"
    summary.freeze_panes = "A5"

    # ---------------- Data ----------------
    descs = [d for d, _, _ in C.QSAR_DESCRIPTORS if any(r.get(d) not in (None, "") for r in data
                                                         if r["level"] in LEVELS)]
    data_ws.append(["level", "compound", "label", "species", "group", "kobs (1/h)", "SE (1/h)", "ln kobs"] + descs)
    bold_row(data_ws, 1)
    num = lambda v: float(v) if v not in (None, "") else None  # noqa: E731
    for r in data:
        if r["level"] in LEVELS:
            data_ws.append([SHORT[r["level"]], r["compound"], r["label"], r["species"], r["group"],
                            num(r["kobs_per_h"]), num(r.get("se_per_h")), num(r["ln_kobs"])]
                           + [num(r.get(d)) for d in descs])
    data_ws.freeze_panes = "C2"
    for i in range(1, data_ws.max_column + 1):
        data_ws.column_dimensions[get_column_letter(i)].width = 14

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    X.patch_many(out, st, jobs)
    print(f"wrote {out}: Summary, Data and {len(rows)} chart tabs")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "figures/M062X_QSAR.xlsx",
         sys.argv[2] if len(sys.argv) > 2 else None)
