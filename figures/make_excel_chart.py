#!/usr/bin/env python3
"""Build an Excel workbook with a scatter chart in Jonas's figure style.

Copied from the plot-like-jonas skill (scripts/make_excel_chart.py). Changes: add_chart_sheet()
and patch_many() build many chart tabs in one workbook in a single pass (re-opening a
workbook with openpyxl does not keep charts reliably), and a filled_gray_square preset.

Usage:
    python make_excel_chart.py spec.json

The data is written to cells and the chart references those cells, so Jonas can
edit numbers in Excel and the chart updates. See references/excel.md for the
spec format. Style defaults live in STYLE below; a spec's "style" key overrides.
"""
import json
import re
import shutil
import sys
import tempfile
import zipfile
from copy import deepcopy

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend, LegendEntry
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Font, Paragraph,
                                   ParagraphProperties, RegularTextRun,
                                   RichTextProperties)
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font as CellFont
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

STYLE = {
    "font": "Cambria",
    "tick_pt": 10, "tick_bold": False,
    "axis_title_pt": 10, "axis_title_bold": False,
    "legend_pt": 10, "legend_bold": False,
    "label_pt": 9, "label_bold": False,
    "width_in": 3.5, "height_in": 3.0,
    "axis_line_pt": 1.0,
    "marker_size": 7,
    "fit_line_pt": 1.5,
    "leader_color": "A6A6A6",
}

# Marker presets: (symbol, fill, edge). Pick by the series' role (see SKILL.md):
#   gray_circle         - default / single data series (gray fill, black outline)
#   open_black_circle   - "new data" when contrasted with published data
#   filled_gray_circle  - published or literature comparison data (no outline)
MARKERS = {
    "gray_circle":         ("circle",   "A6A6A6", "000000"),
    "open_black_circle":   ("circle",   None,     "000000"),
    "filled_gray_circle":  ("circle",   "A6A6A6", "A6A6A6"),
    "filled_black_circle": ("circle",   "000000", "000000"),
    "gray_square":         ("square",   "A6A6A6", "000000"),
    "filled_gray_square":  ("square",   "A6A6A6", "A6A6A6"),   # published data, extra series
    "open_black_square":   ("square",   None,     "000000"),
    "gray_triangle":       ("triangle", "A6A6A6", "000000"),
    "open_black_triangle": ("triangle", None,     "000000"),
}
DEFAULT_ORDER = list(MARKERS)

EMU_PER_PT = 12700
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C15 = "http://schemas.microsoft.com/office/drawing/2012/chart"
NS = {"c": C, "a": A}


def char_props(st, size, bold, baseline=None):
    cp = CharacterProperties(latin=Font(typeface=st["font"]), sz=int(size * 100),
                             b=bold, solidFill="000000")
    if baseline:
        cp.baseline = baseline
    return cp


def text_props(st, size, bold):
    """txPr for tick labels / legend."""
    cp = char_props(st, size, bold)
    return RichText(bodyPr=RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def parse_markup(s):
    """'ln k_{obs}' -> [('ln k', None), ('obs', -25000)]. Supports _{} and ^{}."""
    out, pos = [], 0
    for m in re.finditer(r"([_^])\{([^}]*)\}", s):
        if m.start() > pos:
            out.append((s[pos:m.start()], None))
        out.append((m.group(2), -25000 if m.group(1) == "_" else 30000))
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], None))
    return out


def axis_title(st, text, vertical=False):
    size, bold = st["axis_title_pt"], st["axis_title_bold"]
    runs = [RegularTextRun(rPr=char_props(st, size, bold, base), t=t)
            for t, base in parse_markup(text)]
    body = RichTextProperties(rot=-5400000, vert="horz") if vertical else RichTextProperties()
    para = Paragraph(pPr=ParagraphProperties(defRPr=char_props(st, size, bold)), r=runs)
    return Title(tx=Text(rich=RichText(bodyPr=body, p=[para])), overlay=False)


def line(color="000000", pt=1.0, dash=None):
    ln = LineProperties(solidFill=color, w=int(pt * EMU_PER_PT))
    if dash:
        ln.prstDash = dash
    return ln


def style_axis(ax, st, cfg, title, vertical):
    ax.delete = False
    ax.majorGridlines = None
    ax.minorGridlines = None
    ax.majorTickMark = "out"
    ax.minorTickMark = "none"
    ax.tickLblPos = "low"
    ax.crosses = "min"
    if "min" in cfg: ax.scaling.min = cfg["min"]
    if "max" in cfg: ax.scaling.max = cfg["max"]
    if "major" in cfg: ax.majorUnit = cfg["major"]
    ax.number_format = cfg.get("format", "0.00")
    ax.numFmt.sourceLinked = False
    ax.graphicalProperties = GraphicalProperties(ln=line(pt=st["axis_line_pt"]))
    ax.txPr = text_props(st, st["tick_pt"], st["tick_bold"])
    if title:
        ax.title = axis_title(st, title, vertical)


def add_chart_sheet(ws, spec, st):
    """Data, fit formulas, chart and caption for one spec on worksheet ws. Returns the label jobs
    for patch_chart()."""
    chart = ScatterChart()
    chart.style = None
    chart.scatterStyle = "lineMarker"
    chart.title = None
    chart.width = st["width_in"] * 2.54
    chart.height = st["height_in"] * 2.54

    col = 1
    label_jobs = []  # (series_index, labels, offsets)
    fit_jobs = []  # (series_index, x_col, n)
    n_trend = 0
    for i, s in enumerate(spec["series"]):
        n = len(s["x"])
        ws.cell(row=1, column=col, value=s.get("x_header", spec.get("x_title", "x")))
        ws.cell(row=1, column=col + 1, value=s["name"])
        for c in (col, col + 1):
            ws.cell(row=1, column=c).font = CellFont(bold=True)
        extra = 2
        for r in range(n):
            ws.cell(row=r + 2, column=col, value=s["x"][r])
            ws.cell(row=r + 2, column=col + 1, value=s["y"][r])
        if s.get("labels"):
            ws.cell(row=1, column=col + extra, value="label").font = CellFont(bold=True)
            for r, lab in enumerate(s["labels"]):
                ws.cell(row=r + 2, column=col + extra, value=lab)
            extra += 1
        err_col = None
        if s.get("y_err"):
            err_col = col + extra
            ws.cell(row=1, column=err_col, value="y error").font = CellFont(bold=True)
            for r, e in enumerate(s["y_err"]):
                ws.cell(row=r + 2, column=err_col, value=e)
            extra += 1

        xref = Reference(ws, min_col=col, min_row=2, max_row=n + 1)
        yref = Reference(ws, min_col=col + 1, min_row=1, max_row=n + 1)
        ser = Series(yref, xref, title_from_data=True)

        preset = s.get("style") or DEFAULT_ORDER[i % len(DEFAULT_ORDER)]
        symbol, fill, edge = MARKERS[preset]
        ser.marker = Marker(symbol=symbol, size=s.get("marker_size", st["marker_size"]))
        gp = GraphicalProperties(ln=line(edge, 0.75))
        if fill:
            gp.solidFill = fill
        else:
            gp.noFill = True
        ser.marker.graphicalProperties = gp
        ser.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        ser.smooth = False

        if s.get("trendline"):
            fit_jobs.append((i, col, n))
            extra += 1  # spacer; fit block is written after the loop

        if err_col:
            ref = f"'{ws.title}'!${chr(64 + err_col)}$2:${chr(64 + err_col)}${n + 1}"
            src = NumDataSource(numRef=NumRef(f=ref))
            ser.errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                    noEndCap=False, plus=src, minus=src,
                                    spPr=GraphicalProperties(ln=line(pt=0.75)))

        chart.series.append(ser)
        if s.get("labels"):
            label_jobs.append((i, s["labels"], s.get("label_offsets"), s.get("label_positions")))
        col += extra + 1  # blank spacer column between series

    # Fit lines: a separate two-point line series per fitted data series, driven by
    # SLOPE/INTERCEPT formulas so it updates with the data. Using a series instead
    # of Excel's native trendline keeps legend indexing identical in Excel and
    # LibreOffice (fit series come last, so their legend entries are easy to hide).
    for i, xc, n in fit_jobs:
        s = spec["series"][i]
        X = f"${get_column_letter(xc)}$2:${get_column_letter(xc)}${n + 1}"
        Y = f"${get_column_letter(xc + 1)}$2:${get_column_letter(xc + 1)}${n + 1}"
        fx, fy = get_column_letter(col), get_column_letter(col + 1)
        ws.cell(row=1, column=col, value="fit x").font = CellFont(bold=True)
        ws.cell(row=1, column=col + 1, value=f"fit: {s['name']}").font = CellFont(bold=True)
        ws.cell(row=2, column=col, value=f"=MIN({X})")
        ws.cell(row=3, column=col, value=f"=MAX({X})")
        for r in (2, 3):
            ws.cell(row=r, column=col + 1, value=f"=${fy}$6+${fy}$5*{fx}{r}")
        for r, (lab, f) in enumerate([("slope", f"=SLOPE({Y},{X})"),
                                      ("intercept", f"=INTERCEPT({Y},{X})"),
                                      ("R2", f"=RSQ({Y},{X})")], start=5):
            ws.cell(row=r, column=col, value=lab).font = CellFont(bold=True)
            ws.cell(row=r, column=col + 1, value=f)
        fit = Series(Reference(ws, min_col=col + 1, min_row=1, max_row=3),
                     Reference(ws, min_col=col, min_row=2, max_row=3), title_from_data=True)
        fit.marker = Marker(symbol="none")
        fit.graphicalProperties = GraphicalProperties(
            ln=line(pt=st["fit_line_pt"], dash="sysDot"))
        fit.smooth = False
        chart.series.append(fit)
        n_trend += 1
        col += 3

    style_axis(chart.x_axis, st, spec.get("x_axis", {}), spec.get("x_title"), False)
    style_axis(chart.y_axis, st, spec.get("y_axis", {}), spec.get("y_title"), True)

    # White background everywhere; black box around the plot area; no chart border.
    chart.plot_area.graphicalProperties = GraphicalProperties(
        solidFill="FFFFFF", ln=line(pt=st["axis_line_pt"]))
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True))

    leg = spec.get("legend", {"x": 0.05, "y": 0.72})
    if leg is False or len(spec["series"]) < 2 and not spec.get("legend"):
        chart.legend = None
    else:
        chart.legend = Legend()
        chart.legend.overlay = True
        chart.legend.txPr = text_props(st, st["legend_pt"], st["legend_bold"])
        ml = ManualLayout(xMode="edge", yMode="edge",
                          x=leg.get("x", 0.05), y=leg.get("y", 0.72))
        if "w" in leg: ml.w = leg["w"]
        if "h" in leg: ml.h = leg["h"]
        chart.legend.layout = Layout(manualLayout=ml)
        # Hide the legend entries of the fit-line series (always last).
        ns = len(spec["series"])
        chart.legend.legendEntry = [LegendEntry(idx=ns + k, delete=True) for k in range(n_trend)]

    for c in range(1, col):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.page_setup.orientation = "landscape"
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1

    anchor = spec.get("anchor", f"{get_column_letter(col)}2")
    ws.add_chart(chart, anchor)
    a_col, a_row = coordinate_to_tuple(anchor)[1], coordinate_to_tuple(anchor)[0]

    if spec.get("inset_image"):
        add_inset(ws, spec["inset_image"], st, a_col, a_row)
    if spec.get("caption"):
        add_caption(ws, spec["caption"], st, a_col, a_row)

    return label_jobs


def build(spec):
    st = {**STYLE, **spec.get("style", {})}
    out = spec["output"]
    if spec.get("append_to"):
        wb = load_workbook(spec["append_to"])
        ws = wb.create_sheet(spec.get("sheet", "Figure"))
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = spec.get("sheet", "Figure")

    label_jobs = add_chart_sheet(ws, spec, st)

    # Formulas (fit, slope, R2) have no cached values from openpyxl; Excel computes
    # them on open. Don't run LibreOffice recalc on the deliverable: its re-save
    # drops chart styling (leader lines, open markers). Preview a copy instead.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    patch_file(out, st, label_jobs)
    return out


# Default Excel grid: 64 px columns, 20 px rows (at 96 dpi). Used to place insets
# over the chart, since the chart and image are both anchored to sheet cells.
COL_PX, ROW_PX, EMU_PER_PX = 64, 20, 9525


def sheet_anchor(a_col, a_row, x_px, y_px):
    c, cx = divmod(int(x_px), COL_PX)
    r, ry = divmod(int(y_px), ROW_PX)
    return AnchorMarker(col=a_col - 1 + c, colOff=cx * EMU_PER_PX,
                        row=a_row - 1 + r, rowOff=ry * EMU_PER_PX)


def add_inset(ws, cfg, st, a_col, a_row):
    """Float an image (e.g., a chemical structure PNG) over the chart.

    cfg: {"path": "tcb.png", "x": 0.62, "y": 0.06, "width_in": 0.9}
    x, y are the image's top-left corner as fractions of the chart size.
    """
    img = XLImage(cfg["path"])
    w_px = cfg.get("width_in", 0.9) * 96
    img.height = img.height * w_px / img.width
    img.width = w_px
    x = cfg.get("x", 0.62) * st["width_in"] * 96
    y = cfg.get("y", 0.06) * st["height_in"] * 96
    img.anchor = OneCellAnchor(_from=sheet_anchor(a_col, a_row, x, y),
                               ext=XDRPositiveSize2D(int(img.width * EMU_PER_PX),
                                                     int(img.height * EMU_PER_PX)))
    ws.add_image(img)


def add_caption(ws, text, st, a_col, a_row):
    """Caption in Cambria in merged cells directly under the chart."""
    ncols = max(1, round(st["width_in"] * 96 / COL_PX))
    top = a_row + int(st["height_in"] * 96 / ROW_PX) + 1
    lines = max(2, len(text) // (ncols * 9) + 1)
    ws.merge_cells(start_row=top, start_column=a_col, end_row=top + lines - 1,
                   end_column=a_col + ncols - 1)
    cell = ws.cell(row=top, column=a_col, value=text)
    cell.font = CellFont(name=st["font"], size=st.get("caption_pt", 10))
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def patch_file(path, st, jobs):
    """Post-process the newest chart's XML for things openpyxl can't write:
    autoTitleDeleted (no "Chart Title" placeholder) and per-point custom-text
    data labels with gray leader lines (c15 extension)."""
    tmp = tempfile.mktemp(suffix=".xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        chart_name = max((n for n in zin.namelist() if re.match(r"xl/charts/chart\d+\.xml", n)), key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == chart_name:
                data = patch_chart(data, st, jobs)
            zout.writestr(item, data)
    shutil.move(tmp, path)


def patch_many(path, st, jobs_by_sheet):
    """Like patch_file, for every chart in the workbook: each chart is matched to its sheet
    through the sheet name in its first series formula."""
    tmp = tempfile.mktemp(suffix=".xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r"xl/charts/chart\d+\.xml", item.filename):
                m = re.search(rb"<(?:c:)?f>'?([^'!<]+)'?!", data)   # openpyxl writes <f>
                sheet = m.group(1).decode() if m else None
                data = patch_chart(data, st, jobs_by_sheet.get(sheet, []))
            zout.writestr(item, data)
    shutil.move(tmp, path)


def E(tag, **attrs):
    p, t = tag.split(":")
    el = etree.Element(f"{{{NS[p]}}}{t}")
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def sub(parent, tag, **attrs):
    el = E(tag, **attrs)
    parent.append(el)
    return el


def rpr(parent, tag, st):
    r = sub(parent, tag, lang="en-US", sz=int(st["label_pt"] * 100),
            b="1" if st["label_bold"] else "0")
    sf = sub(r, "a:solidFill"); sub(sf, "a:srgbClr", val="000000")
    sub(r, "a:latin", typeface=st["font"])
    return r


def flags(parent):
    for tag, v in [("showLegendKey", 0), ("showVal", 1), ("showCatName", 0),
                   ("showSerName", 0), ("showPercent", 0), ("showBubbleSize", 0)]:
        sub(parent, f"c:{tag}", val=v)


def patch_chart(xml, st, jobs):
    root = etree.fromstring(xml)
    chart_el = root.find("c:chart", NS)
    atd = chart_el.find("c:autoTitleDeleted", NS)
    if atd is None:
        atd = E("c:autoTitleDeleted")
        title = chart_el.find("c:title", NS)
        (title.addnext if title is not None else chart_el.insert)(*(
            [atd] if title is not None else [0, atd]))
    atd.set("val", "1")
    sers = root.findall(".//c:scatterChart/c:ser", NS)
    for idx, labels, offsets, positions in jobs:
        ser = sers[idx]
        for old in ser.findall("c:dLbls", NS):
            ser.remove(old)
        dl = E("c:dLbls")
        for k, text in enumerate(labels):
            if text in (None, ""):
                d = sub(dl, "c:dLbl"); sub(d, "c:idx", val=k); sub(d, "c:delete", val=1)
                continue
            d = sub(dl, "c:dLbl")
            sub(d, "c:idx", val=k)
            if offsets and offsets[k]:
                lay = sub(d, "c:layout"); ml = sub(lay, "c:manualLayout")
                sub(ml, "c:x", val=offsets[k][0]); sub(ml, "c:y", val=offsets[k][1])
            tx = sub(d, "c:tx"); rich = sub(tx, "c:rich")
            sub(rich, "a:bodyPr"); p = sub(rich, "a:p")
            ppr = sub(p, "a:pPr"); rpr(ppr, "a:defRPr", st)
            r = sub(p, "a:r"); rpr(r, "a:rPr", st); sub(r, "a:t").text = text
            pos = (positions[k] if positions and positions[k] else "r")
            sub(d, "c:dLblPos", val={"left": "l", "right": "r", "above": "t",
                                     "below": "b"}.get(pos, pos))
            flags(d)
        sp = sub(dl, "c:spPr"); sub(sp, "a:noFill"); ln = sub(sp, "a:ln"); sub(ln, "a:noFill")
        sub(dl, "c:dLblPos", val="r")
        flags(dl)
        ext_lst = sub(dl, "c:extLst")
        ext = sub(ext_lst, "c:ext", uri="{CE6537A1-D6FC-4f65-9D91-7224C49458BB}")
        sll = etree.SubElement(ext, f"{{{C15}}}showLeaderLines", nsmap={"c15": C15}); sll.set("val", "1")
        ll = etree.SubElement(ext, f"{{{C15}}}leaderLines", nsmap={"c15": C15})
        llsp = etree.SubElement(ll, f"{{{C}}}spPr"); lln = sub(llsp, "a:ln", w=int(0.75 * EMU_PER_PT))
        sf = sub(lln, "a:solidFill"); sub(sf, "a:srgbClr", val=st["leader_color"])
        # dLbls goes after marker/dPt and before trendline/errBars/xVal.
        anchor = None
        for tag in ("trendline", "errBars", "xVal", "yVal"):
            anchor = ser.find(f"c:{tag}", NS)
            if anchor is not None:
                break
        anchor.addprevious(dl)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


if __name__ == "__main__":
    with open(sys.argv[1]) as f:
        print(build(json.load(f)))
