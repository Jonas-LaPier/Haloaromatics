#!/usr/bin/env python3
"""Haloaromatics workflow driver.

    python3 hx.py species                     # list species / reactions  -> results/
    python3 hx.py generate am1                # write .gjf inputs
    python3 hx.py submit am1 [--dry-run]      # job array of every input without a log
    python3 hx.py submit am1 --serial --time 0-00:30:00   # or all of them in one job
    python3 hx.py status am1                  # QC summary
    python3 hx.py retry am1                   # rebuild failed jobs with fixes, then submit again
    python3 hx.py scrape all                  # logs -> results/raw/<stage>.csv
    python3 hx.py compile                     # raw csv -> reactions, barriers, xlsx

Stage names: am1, b3lyp_gas, b3lyp_smd, m062x_gas, m062x_smd,
             tsscan_m062x_gas, tsscan_m062x_smd, ts_m062x_gas, ts_m062x_smd
Groups:      optfreq (4 DFT levels), tsscan, ts, all
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import shutil
import subprocess
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))

import config as C  # noqa: E402
from haloaro import correlations, descriptors, qc, thermo  # noqa: E402
from haloaro.geomtools import sphere_radius  # noqa: E402
from haloaro.parse import parse_log  # noqa: E402
from haloaro.stages import (REACTIONS, SPECIES, STAGES, Skip, build_input, jobs, resources,  # noqa: E402
                            resolve, retry_fixes)

RESULTS = Path("results")
RAW = RESULTS / "raw"


# --------------------------------------------------------------------------- #
def write_csv(path, rows, fields=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("")
        return
    fields = fields or _union_keys(rows)
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def _union_keys(rows):
    keys = {}
    for r in rows:
        for k in r:
            keys.setdefault(k, None)
    return list(keys)


def read_csv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def job_state(stage, job, sp, site):
    lp = stage.log(job)
    p = parse_log(lp) if lp.exists() else None
    status, flags = qc.evaluate(p, sp, stage.kind, site)
    return p, status, flags


# --------------------------------------------------------------------------- #
def cmd_species(a):
    rows = [{"name": s.name, "kind": s.kind, "halogen": s.halogen or "", "n_X": s.n_hal,
             "charge": s.charge, "mult": s.mult,
             "ring_C1..C6": "-".join(s.ring) if s.ring else s.element} for s in SPECIES.values()]
    write_csv(RESULTS / "species_list.csv", rows)
    write_csv(RESULTS / "reaction_list.csv", [vars(r) for r in REACTIONS])
    print(Counter(s.kind for s in SPECIES.values()))
    print(f"{len(REACTIONS)} symmetry-unique dehalogenation reactions")
    for st in STAGES.values():
        print(f"  {st.name:<20} {len(jobs(st)):>4} jobs")
    print(f"wrote {RESULTS/'species_list.csv'} and {RESULTS/'reaction_list.csv'}")


def cmd_generate(a):
    for sn in resolve(a.stages):
        st = STAGES[sn]
        for sub in ("inputs", "logs", "chks"):
            (st.dir / sub).mkdir(parents=True, exist_ok=True)
        n_new = n_exist = 0
        skipped = []
        for job, sp, site in jobs(st):
            if a.only and job not in a.only:
                continue
            ip = st.inp(job)
            if ip.exists() and not a.force:
                n_exist += 1
                continue
            try:
                txt = build_input(st, job, sp, site, ignore_prereq=a.ignore_prereq)
            except Skip as e:
                skipped.append((job, str(e)))
                continue
            ip.write_text(txt)
            n_new += 1
        print(f"[{sn}] {n_new} written, {n_exist} already present, {len(skipped)} skipped")
        show = skipped if a.verbose else skipped[:5]
        for j, why in show:
            print(f"    skip {j}: {why}")
        if len(skipped) > len(show):
            print(f"    ... {len(skipped) - len(show)} more (use --verbose)")


def cmd_submit(a):
    if a.serial and not a.time:
        raise SystemExit("--serial needs an explicit --time for the whole batch (e.g. --time 0-00:30:00)")
    Path("slurm_logs").mkdir(exist_ok=True)
    for sn in resolve(a.stages):
        st = STAGES[sn]
        todo = []
        for job, sp, site in jobs(st):
            if a.only and job not in a.only:
                continue
            ip = st.inp(job)
            if not ip.exists():
                continue
            lp = st.log(job)
            if lp.exists():
                if not a.include_incomplete:
                    continue
                _, status, _ = job_state(st, job, sp, site)
                if status != "incomplete":
                    continue
            todo.append((str(ip), job))
        if not todo:
            print(f"[{sn}] nothing to submit")
            continue
        # one array per resource class (e.g. benzenes vs. diphenyl ethers)
        groups = {}
        for ip, job in todo:
            r = resources(st, job)
            groups.setdefault((r["cpus"], r["mem_gb"], r["time"]), []).append(ip)
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        for gi, ((cpus, mem, time), ips) in enumerate(sorted(groups.items())):
            suffix = "" if len(groups) == 1 else f"_{gi + 1}"
            man = st.dir / f"manifest_{stamp}{suffix}.txt"
            if not a.dry_run:
                man.write_text("\n".join(ips) + "\n")
            cmd = ["sbatch", "--parsable", f"--job-name={sn}"]
            if a.serial:
                # one job running every input in turn (Sherlock: aggregate short tasks)
                cmd += ["--output=slurm_logs/%x.%j.out", "--error=slurm_logs/%x.%j.err"]
            else:
                cmd.append(f"--array=1-{len(ips)}%{C.ARRAY_THROTTLE}")
            cmd += [f"--cpus-per-task={cpus}", f"--mem={mem}G", f"--time={a.time or time}"]
            if C.PARTITION:
                cmd.append(f"--partition={C.PARTITION}")
            cmd += ["slurm/g16_array.sbatch", str(man)]
            print(f"[{sn}] {len(ips)} jobs -> {man}" + (" (serial)" if a.serial else ""))
            print("   " + " ".join(cmd))
            if not a.dry_run:
                if shutil.which("sbatch") is None:
                    print("   sbatch not found (not on Sherlock?) - run the command above manually")
                else:
                    r = subprocess.run(cmd, check=False, capture_output=True, text=True)
                    if r.returncode == 0:
                        print(f"   submitted job {r.stdout.strip().split(';')[0]}")
                    else:
                        print(f"   sbatch failed: {r.stderr.strip()}")


def cmd_status(a):
    for sn in resolve(a.stages):
        st = STAGES[sn]
        cnt = Counter()
        bad = []
        for job, sp, site in jobs(st):
            if not st.inp(job).exists():
                cnt["no_input"] += 1
                continue
            _, status, flags = job_state(st, job, sp, site)
            cnt[status] += 1
            if status not in ("ok",) and (a.verbose or status != "missing"):
                bad.append((job, status, flags))
        print(f"[{sn}] " + ", ".join(f"{k}={v}" for k, v in sorted(cnt.items())))
        for job, status, flags in bad:
            print(f"    {status:<10} {job:<28} {'; '.join(flags)}")


def cmd_retry(a):
    for sn in resolve(a.stages):
        st = STAGES[sn]
        (st.dir / "logs" / "failed").mkdir(parents=True, exist_ok=True)
        for job, sp, site in jobs(st):
            if a.only and job not in a.only:
                continue
            if not st.log(job).exists():
                continue
            p, status, flags = job_state(st, job, sp, site)
            if status not in ("fail",) + (("incomplete",) if a.include_incomplete else ()):
                continue
            n_prev = len(list((st.dir / "logs" / "failed").glob(f"{job}.try*.log")))
            if n_prev >= C.MAX_RETRIES and not a.force:
                print(f"[{sn}] {job}: already retried {n_prev}x ({flags}); inspect manually or use --force")
                continue
            fx = retry_fixes(st, p, status, flags, sp)
            if not fx and not a.force:
                print(f"[{sn}] {job}: no automatic fix for {flags}; inspect manually (or --force to rerun as-is)")
                continue
            n = len(list((st.dir / "logs" / "failed").glob(f"{job}.try*.log"))) + 1
            shutil.move(str(st.log(job)), st.dir / "logs" / "failed" / f"{job}.try{n}.log")
            shutil.copy(st.inp(job), st.dir / "logs" / "failed" / f"{job}.try{n}.gjf")
            try:
                st.inp(job).write_text(build_input(st, job, sp, site, retry=fx, ignore_prereq=True))
            except Skip as e:
                print(f"[{sn}] {job}: could not rebuild ({e})")
                continue
            print(f"[{sn}] {job}: {status} {flags} -> retry #{n} with {sorted(fx)}")
    print("Now run: python3 hx.py submit <stage>")


def _ra_cols(sp, p, st):
    """Radical-anion character (pi / sigma_bent / dissociated / elongated) from its geometry."""
    if sp.kind != "radical_anion" or st.kind != "optfreq" or not p.get("geometry"):
        return {}
    state, site, r, oop = qc.ra_state(p["geometry"], sp)
    return {"RA_state": state, "RA_longest_CX_site": site,
            "RA_longest_CX_A": round(r, 3) if r else None, "RA_oop_deg": round(oop, 1) if oop is not None else None}


def cmd_scrape(a):
    for sn in resolve(a.stages):
        st = STAGES[sn]
        rows = []
        atoms = {}
        for job, sp, site in jobs(st):
            if not st.inp(job).exists():
                continue
            p, status, flags = job_state(st, job, sp, site)
            p = p or {}
            if st.kind != "am1" and p.get("geometry"):
                atoms[job] = {"status": status, "geometry": p["geometry"],
                              "hirshfeld": p.get("hirshfeld"), "wiberg": p.get("wiberg")}
            lumo = p.get("lumo")
            rows.append({
                "stage": sn, "level": st.level or "am1", "name": job, "kind": sp.kind,
                "halogen": sp.halogen or "", "n_X": sp.n_hal, "charge": sp.charge, "mult": sp.mult,
                "status": status, "flags": "; ".join(flags),
                "E_scf": p.get("E_scf"), "G": p.get("G"), "H": p.get("H"), "ZPE": p.get("ZPE"),
                "G_corr": p.get("G_corr"), "homo": p.get("homo"), "lumo": lumo,
                "lumo_eV": round(lumo * C.HARTREE_TO_EV, 4) if lumo is not None else None,
                "S2": p.get("S2"), "n_imag": p.get("n_imag"), "lowest_freq": p.get("lowest_freq"),
                "wall_hours": p.get("wall_hours"),
                "dipole_D": p.get("dipole_D"),
                "polar_iso_A3": (round(p["polar_iso_bohr3"] * 0.148185, 3)
                                 if p.get("polar_iso_bohr3") is not None else None),
                **_ra_cols(sp, p, st),
                "radius_A": (round(sphere_radius(p["geometry"]), 3)
                             if sp.kind == "parent" and p.get("geometry") and qc.usable(status) else None),
            })
        write_csv(RAW / f"{sn}.csv", rows)
        if atoms:
            (RAW / f"{sn}_atoms.json").write_text(json.dumps(atoms))
        c = Counter(r["status"] for r in rows)
        print(f"[{sn}] {len(rows)} rows -> {RAW / (sn + '.csv')}  {dict(c)}")


def cmd_compile(a):
    tab = {}
    species_rows = []
    for sn, st in STAGES.items():
        f = RAW / f"{sn}.csv"
        if not f.exists() or f.stat().st_size == 0:
            continue
        for r in read_csv(f):
            key_level = sn if st.kind in ("ts", "tsscan") else r["level"]
            tab[(key_level, r["name"])] = r
            species_rows.append(r)
    if not tab:
        raise SystemExit("No scraped data. Run: python3 hx.py scrape all")

    rx = thermo.reaction_table(tab, REACTIONS, SPECIES)
    ts = thermo.ts_table(tab, REACTIONS) if C.RUN_RA_TS else []
    det = thermo.det_table(tab, REACTIONS)
    atoms = {}
    for sn, st in STAGES.items():
        f = RAW / f"{sn}_atoms.json"
        if f.exists() and st.kind in ("optfreq", "sp"):
            for name, v in json.loads(f.read_text()).items():
                atoms[(st.level, name)] = v
    mol, lam_i = descriptors.molecular_table(tab, SPECIES, rx)
    path = thermo.pathway_table(rx, ts, det, lam_i) if C.RUN_RA_TS else []
    sites = descriptors.site_table(tab, atoms, SPECIES, REACTIONS, rx, ts, det)
    prod = descriptors.product_distribution(sites, det, path)
    corr, cdata, cpairs, cpred = correlations.run(mol, rx, det, path, ts, sites)
    lumo = [{"level": r["level"], "name": r["name"], "halogen": r["halogen"], "n_X": r["n_X"],
             "LUMO_Eh": r["lumo"], "LUMO_eV": r["lumo_eV"], "HOMO_Eh": r["homo"], "status": r["status"]}
            for r in species_rows if r["kind"] == "parent" and r["level"] in C.LEVELS]
    qcrows = [r for r in species_rows if r["status"] != "ok"]
    consts = [{"constant": k, "value": getattr(C, k)} for k in (
        "BASIS", "SOLVENT", "TEMPERATURE", "STD_STATE_CORR_KCAL", "G_ELECTRON_KCAL", "E_ABS_SHE_V",
        "G_PROTON_GAS_KCAL", "DG_SOLV_PROTON_KCAL", "HARTREE_TO_KCAL", "FARADAY_KCAL",
        "SAVEANT_D", "SPIN_ORBIT_KCAL", "EPS_STATIC", "EPS_OPTICAL", "LAMBDA0_MODEL",
        "LAMBDA0_KCAL", "RADIUS_PROBE_A", "DET_POTENTIALS_V")]

    write_csv(RESULTS / "species_energies.csv", species_rows)
    write_csv(RESULTS / "parent_LUMO.csv", lumo)
    write_csv(RESULTS / "reactions.csv", rx)
    write_csv(RESULTS / "det_barriers.csv", det)
    write_csv(RESULTS / "molecular_descriptors.csv", mol)
    write_csv(RESULTS / "site_descriptors.csv", sites)
    write_csv(RESULTS / "product_distribution.csv", prod)
    if corr:
        write_csv(RESULTS / "correlations.csv", corr)
        write_csv(RESULTS / "correlation_data.csv", cdata)
        write_csv(RESULTS / "correlation_pairs.csv", cpairs)
        write_csv(RESULTS / "correlation_predictions.csv", cpred)
        figs = correlations.plots(corr, cdata, RESULTS / "plots")
        print(f"wrote results/correlations*.csv ({len(corr)} fits)" + (f" and {len(figs)} plots in results/plots/" if figs else ""))
    if C.RUN_RA_TS:
        write_csv(RESULTS / "ts_barriers.csv", ts)
        write_csv(RESULTS / "pathway_comparison.csv", path)
    write_csv(RESULTS / "qc_issues.csv", qcrows)
    print(f"wrote results/species_energies.csv ({len(species_rows)}), parent_LUMO.csv ({len(lumo)}), "
          f"reactions.csv ({len(rx)}), det_barriers.csv ({len(det)}), molecular_descriptors.csv ({len(mol)}), "
          f"site_descriptors.csv ({len(sites)}), product_distribution.csv ({len(prod)}), "
          f"qc_issues.csv ({len(qcrows)})")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl not installed: skipping xlsx (pip install --user openpyxl)")
        return
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [("README", [{"note": thermo.__doc__ + (thermo.pathway_table.__doc__ if C.RUN_RA_TS else "")
                                     + descriptors.__doc__ + correlations.__doc__}]),
              ("Molecular_descriptors", mol), ("Site_descriptors", sites),
              ("Product_distribution", prod), ("Correlations", corr), ("Correlation_data", cdata),
              ("Correlation_pairs", cpairs), ("Correlation_predictions", cpred),
              ("Reactions", rx), ("DET_barriers", det)]
    if C.RUN_RA_TS:
        sheets += [("TS_barriers", ts), ("Pathway_comparison", path)]
    sheets += [
              ("Parent_LUMO", lumo), ("Species", species_rows), ("QC_issues", qcrows),
              ("Constants", consts)]
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        if not rows:
            continue
        hdr = _union_keys(rows)
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_num(r.get(h)) if not isinstance(r.get(h), (list, dict, tuple)) else str(r.get(h))
                       for h in hdr])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                60, max(10, max(len(str(c.value or "")) for c in col[:50]) + 2))
    out = RESULTS / "Haloaromatics_results.xlsx"
    wb.save(out)
    print(f"wrote {out}")


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(v) if any(ch in v for ch in ".eE") else int(v)
    except (ValueError, TypeError):
        return v


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("species").set_defaults(f=cmd_species)

    def stage_cmd(name, f, **flags):
        p = sub.add_parser(name)
        p.add_argument("stages", nargs="+")
        p.add_argument("--only", nargs="*", help="restrict to these job names")
        for fl, h in flags.items():
            p.add_argument(f"--{fl.replace('_', '-')}", action="store_true", help=h)
        p.set_defaults(f=f)
        return p

    stage_cmd("generate", cmd_generate, force="overwrite existing inputs",
              ignore_prereq="do not require the AM1 log to pass QC (opt/freq stages only)",
              verbose="list every skipped job")
    sp_submit = stage_cmd("submit", cmd_submit, dry_run="print sbatch command only",
                          include_incomplete="resubmit logs without a termination line (make sure they are not running!)",
                          serial="run all inputs one after another in a single job (for short jobs such as AM1)")
    sp_submit.add_argument("--time", help="override the Slurm time limit, e.g. 0-00:30:00 (needed with --serial)")
    stage_cmd("status", cmd_status, verbose="also list missing jobs")
    stage_cmd("retry", cmd_retry, force="rerun failures without an automatic fix",
              include_incomplete="also retry jobs that were killed (e.g. time limit)")
    stage_cmd("scrape", cmd_scrape)
    sub.add_parser("compile").set_defaults(f=cmd_compile)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
