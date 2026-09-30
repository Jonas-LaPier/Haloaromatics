"""Multiple-linear-regression QSAR exploration workbook (separate from the main QSAR workbook).

Run from the project root after compile:
    .venv/bin/python figures/qsar_mlr_explore.py figures/QSAR_MLR_exploration.xlsx [preview_dir]

Every combination of 1-3 descriptors at each level is fitted and scored on repeated stratified
train/test splits (haloaro/mlr.py; settings MLR_* in config.py). Tabs:
    Summary          key findings (computed), best models, baselines, descriptor frequency
    Interpretation   findings next to what the literature in Jonas's Zotero library says
                     (paraphrased from passages read in the source; citation and DOI given)
    Methods          splits, statistics, collinearity, caveats
    All models       every model with its scores
    Candidates       descriptors used and excluded at each level
    Parity ...       out-of-sample predicted vs observed ln kobs for four selected models
The Interpretation text is written for the current results; re-read it if the data change.
"""
from __future__ import annotations

import collections
import csv
import sys
from pathlib import Path

sys.path.insert(0, ".")
sys.path.insert(0, "figures")
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402

import config as C  # noqa: E402
import make_excel_chart as X  # noqa: E402
import qsar_workbook as QW  # noqa: E402
from haloaro import mlr  # noqa: E402
from haloaro import plotting as P  # noqa: E402

BOLD, HEAD = Font(bold=True), Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical="top")
SHADE = PatternFill("solid", fgColor="E7E6E6")

REFS = {
    "LaPier2026": "LaPier, J. K.; Liu, Y.-J.; King, J. F.; Béguerie, T.; Nzihou, A.; Mitch, W. A. Electrochemical "
                  "Debromination of Brominated Aromatic Flame Retardants Using Activated Carbon-Based Cathodes. "
                  "Environ. Sci. Technol. 2026, 60, 1346. https://doi.org/10.1021/acs.est.5c03324",
    "KingMitch2022": "King, J. F.; Mitch, W. A. Electrochemical Reduction of Halogenated Alkanes and Alkenes Using "
                     "Activated Carbon-Based Cathodes. Environ. Sci. Technol. 2022, 56, 17965. "
                     "https://doi.org/10.1021/acs.est.2c05608",
    "KingMitch2024": "King, J. F.; Mitch, W. A. Electrochemical reduction of halogenated organic contaminants using "
                     "carbon-based cathodes: A review. Crit. Rev. Environ. Sci. Technol. 2024. "
                     "https://doi.org/10.1080/10643389.2023.2239130",
    "Neukermans2020": "Neukermans, S.; Vorobjov, F.; Kenis, T.; De Wolf, R.; Hereijgers, J.; Breugelmans, T. "
                      "Electrochemical reduction of halogenated aromatic compounds at metal cathodes in acetonitrile. "
                      "Electrochim. Acta 2020. https://doi.org/10.1016/j.electacta.2019.135484",
    "Andrieux1986": "Andrieux, C. P.; Savéant, J. M.; Su, K. B. Kinetics of dissociative electron transfer. Direct "
                    "and mediated electrochemical reductive cleavage of the carbon-halogen bond. J. Phys. Chem. "
                    "1986, 90, 3815. https://doi.org/10.1021/j100407a059",
    "Mazzucato2023": "Mazzucato, M.; Isse, A. A.; Durante, C. Dissociative electron transfer mechanism and "
                     "application in the electrocatalytic activation of organic halides. Curr. Opin. Electrochem. "
                     "2023. https://doi.org/10.1016/j.coelec.2023.101254",
}


def label(d):
    lab = {k: v for k, v, _ in C.QSAR_DESCRIPTORS}.get(d, d)
    return P.chem(lab, plain=True)


def model_name(m):
    return " + ".join(label(d) for d in m["descriptors"])


def write_table(ws, rows, r0):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r0 + i, column=j + 1, value=v)
            c.alignment = WRAP
            if i == 0:
                c.font, c.fill = BOLD, SHADE
            elif isinstance(v, float):
                c.number_format = "0.00"
    return r0 + len(rows) + 1


def para(ws, r, title, text, ncol=8):
    if title:
        ws.cell(row=r, column=1, value=title).font = HEAD
        r += 1
    c = ws.cell(row=r, column=1, value=text)
    c.alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    return r + 1


def parity_sheet(wb, name, data, m, st, jobs, preview_dir):
    """Out-of-sample predicted vs observed chart for model m."""
    pred = mlr.oos_predictions(data, m["level"], m["descriptors"])
    pts = [{"x": yo, "y": yp, "group": r["group"], "label": r.get("label") or r["compound"], "yerr": None}
           for r, yo, yp, n in pred if yp is not None]
    lo, hi = -8, 8
    fig = plt.figure(figsize=QW.CHART_IN)
    pa = QW.PLOT_AREA
    ax = fig.add_axes([pa["x"], 1 - pa["y"] - pa["h"], pa["w"], pa["h"]])
    P._series(ax, pts, 7)
    one = ax.plot([lo, hi], [lo, hi], color="#a6a6a6", lw=0.75, zorder=1)[0]
    for s_ in ("x", "y"):
        P._set_ticks(ax, s_, list(range(lo, hi + 1, 2)), 0)
    placed = P._place_labels(fig, ax, pts, one, fontsize=QW.FONTS["label_pt"])
    if preview_dir:
        fig.savefig(f"{preview_dir}/{name.replace(' ', '_')}.png", dpi=150, facecolor="white")
    plt.close(fig)
    series = []
    for grp in P.ROLES:
        sel = [p for p in pts if p["group"] == grp]
        if not sel:
            continue
        pos, offs = [], []
        for p in sel:
            dx, dy = placed.get(p["label"], (6, 0))
            where, (bx, by) = QW.excel_position(dx, dy)
            pos.append(where)
            ox, oy = (dx - bx) / QW.CHART_PT[0], -(dy - by) / QW.CHART_PT[1]
            offs.append([round(ox, 4), round(oy, 4)] if abs(dx - bx) > 2 or abs(dy - by) > 2 else None)
        series.append({"name": P.ROLES[grp]["label"], "x_header": "observed ln kobs",
                       "x": [round(p["x"], 4) for p in sel], "y": [round(p["y"], 4) for p in sel],
                       "labels": [p["label"] for p in sel], "label_positions": pos, "label_offsets": offs,
                       "style": QW.STYLES[grp]})
    series.append({"name": "1:1", "x_header": "1:1", "x": [lo, hi], "y": [lo, hi], "style": "line"})
    cap = (f"Out-of-sample predicted vs observed ln k_obs (h⁻¹) for the model {model_name(m)} at "
           f"{P.level_text(m['level'])}. Each prediction is the mean over the {C.MLR_SPLITS} random "
           f"splits in which that compound was in the test set (model fitted on the other 10 compounds). "
           f"Gray line: 1:1. Test RMSE {m['RMSE_test']:.2f} ± {m['RMSE_test_sd']:.2f} (ln units), "
           f"Q²_F1 {m['Q2_F1']:.2f}.")
    spec = {"x_title": "observed ln k_{obs} (h^{−1})", "y_title": "predicted ln k_{obs} (h^{−1})",
            "x_axis": {"min": lo, "max": hi, "major": 2, "format": "0"},
            "y_axis": {"min": lo, "max": hi, "major": 2, "format": "0"}, "series": series,
            "legend": dict(QW.LEGEND), "plot_area": QW.PLOT_AREA, "caption": cap}
    ws = wb.create_sheet(name)
    jobs[name] = X.add_chart_sheet(ws, spec, st)
    # the legend entry of the 1:1 line is not needed; patch_chart keeps it, which is harmless


def main(out, preview_dir=None):
    data = list(csv.DictReader(open("results/correlation_data.csv")))
    models, info = mlr.explore(data)
    ok = [m for m in models if not m["collinear"]]
    if preview_dir:
        Path(preview_dir).mkdir(parents=True, exist_ok=True)
    P.use_style()
    st = {**X.STYLE, **QW.FONTS, "width_in": QW.CHART_IN[0], "height_in": QW.CHART_IN[1]}

    best = ok[0]
    best_k = {k: next(m for m in ok if m["k"] == k) for k in range(1, C.MLR_MAX_TERMS + 1)}
    best_lv = {lv: next(m for m in ok if m["level"] == lv) for lv in info}
    transfer = sorted([m for m in ok if m["RMSE_transfer"] is not None and m["k"] <= 2],
                      key=lambda m: m["RMSE_transfer"])
    within = [m for m in ok if m["RMSE_test"] <= best["RMSE_test"] + best["RMSE_test_sd"]]
    freq = collections.Counter(d for m in ok[:30] for d in m["descriptors"])
    base = next(iter(info.values()))["RMSE_test_mean_model"]
    q = [r for r in csv.DictReader(open("results/qsar_summary.csv")) if r["descriptors"] == "min_wiberg_CX_parent"
         and r.get("R2") not in (None, "")]
    wr = {g: [float(r["R2"]) for r in q if r["fit_set"] == g] for g in ("bromobenzene", "chlorobenzene", "halobenzene")}
    nx = min((m for m in models if m["descriptors"] == ["n_X"]), key=lambda m: m["RMSE_test"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    r = para(ws, 1, "Multiple-linear-regression QSARs: which combination of parameters predicts ln k_obs best?",
             f"All {len(models)} models with 1–{C.MLR_MAX_TERMS} descriptors ({len(ok)} without strongly correlated "
             f"descriptor pairs, |r| ≤ {C.MLR_MAX_R}) at four levels of theory, each scored on {C.MLR_SPLITS} "
             "random training/test splits (10 training, 4 test compounds: 2 bromobenzenes, 1 chlorobenzene, 1 PBDE). "
             f"Rate constants measured at {P.FIRST_E}. Lower test RMSE is better (ln units; 1 = factor 2.7 in k_obs). "
             "Methods and definitions on the Methods tab; literature context on the Interpretation tab.")
    findings = [
        f"Best model: {model_name(best)} ({P.level_text(best['level'])}): test RMSE {best['RMSE_test']:.2f} ± "
        f"{best['RMSE_test_sd']:.2f}, Q²_F1 {best['Q2_F1']:.2f}. For comparison: predicting the training mean gives "
        f"{base:.2f}, the number of halogens alone {nx['RMSE_test']:.2f}, and the best single descriptor "
        f"({model_name(best_k[1])}, {P.level_text(best_k[1]['level'])}) {best_k[1]['RMSE_test']:.2f}.",
        f"{len(within)} models score within one standard deviation of the best, so the ranking among the top models "
        "is not decisive; look at which descriptors recur rather than at the single winner.",
        f"The Wiberg C–X bond order of the parent appears in {freq['min_wiberg_CX_parent']} of the top 30 models. "
        f"Within each family it tracks ln k_obs closely (single-descriptor R² {min(wr['bromobenzene']):.2f}–"
        f"{max(wr['bromobenzene']):.2f} for the bromobenzenes and {min(wr['chlorobenzene']):.2f}–"
        f"{max(wr['chlorobenzene']):.2f} for the chlorobenzenes across the four levels, but only "
        f"{min(wr['halobenzene']):.2f}–{max(wr['halobenzene']):.2f} for both together), because chlorobenzene "
        "values are shifted to higher bond orders; it therefore needs a second, "
        "halogen-sensitive term (Savéant intrinsic barrier, C–X bond energy, LUMO or a reduction free energy) to "
        "put both families on one line.",
        f"Adding a second descriptor lowers the test error from {best_k[1]['RMSE_test']:.2f} to "
        f"{best_k[2]['RMSE_test']:.2f}; a third adds little ({best_k[3]['RMSE_test']:.2f}), so two descriptors are "
        "the most that these 14 compounds support.",
        "Models trained on the bromobenzenes only and applied to the chlorobenzenes and PBDEs (transfer test) favour "
        f"different descriptors: the best is {model_name(transfer[0])} ({P.level_text(transfer[0]['level'])}, transfer "
        f"RMSE {transfer[0]['RMSE_transfer']:.2f}). Models containing the Wiberg bond order transfer poorly because "
        "the chlorobenzene bond orders lie outside the bromobenzene range (extrapolation).",
        "Recommendation: to predict new bromo- and chloroaromatics, fit a two-descriptor model on a training set "
        "that contains both halogens; to extrapolate from bromobenzenes alone, use a frontier-orbital or "
        "Fukui-type descriptor, which transfers best.",
    ]
    r = para(ws, r, "Key findings", "• " + findings[0])
    for f in findings[1:]:
        r = para(ws, r, None, "• " + f)
    r += 1
    ws.cell(row=r, column=1, value="Top 10 models (test RMSE, non-collinear)").font = HEAD
    hdr = ["rank", "level", "k", "descriptors", "test RMSE", "± sd", "Q²_F1", "train R²", "R² (all 14)",
           "LOO Q²", "transfer RMSE", "max |r|"]
    rows = [hdr] + [[i + 1, QW.SHORT[m["level"]], m["k"], model_name(m), m["RMSE_test"], m["RMSE_test_sd"],
                     m["Q2_F1"], m["R2_train"], m["R2"], m["Q2_LOO"], m["RMSE_transfer"], m["max_abs_r"]]
                    for i, m in enumerate(ok[:10])]
    r = write_table(ws, rows, r + 1)
    ws.cell(row=r, column=1, value="Best model with 1, 2 and 3 descriptors, and per level").font = HEAD
    rows = [hdr[1:4] + hdr[4:11]]
    for tag, m in list(best_k.items()) + list(best_lv.items()):
        rows.append([QW.SHORT[m["level"]], m["k"], model_name(m), m["RMSE_test"], m["RMSE_test_sd"], m["Q2_F1"],
                     m["R2_train"], m["R2"], m["Q2_LOO"], m["RMSE_transfer"]])
    rows.append(["all", 0, "baseline: training mean", base, None, 0.0, None, None, None, None])
    rows.append([QW.SHORT[nx["level"]], 1, "baseline: number of halogens", nx["RMSE_test"], nx["RMSE_test_sd"],
                 nx["Q2_F1"], nx["R2_train"], nx["R2"], nx["Q2_LOO"], nx["RMSE_transfer"]])
    r = write_table(ws, rows, r + 1)
    ws.cell(row=r, column=1, value="Best transfer models (train on bromobenzenes, predict chlorobenzenes + PBDEs)").font = HEAD
    rows = [["level", "k", "descriptors", "transfer RMSE", "test RMSE", "R² (all 14)"]] + \
        [[QW.SHORT[m["level"]], m["k"], model_name(m), m["RMSE_transfer"], m["RMSE_test"], m["R2"]] for m in transfer[:6]]
    r = write_table(ws, rows, r + 1)
    ws.cell(row=r, column=1, value="How often each descriptor appears in the top 30 models").font = HEAD
    rows = [["descriptor", "key", "count"]] + [[label(d), d, n] for d, n in freq.most_common()]
    write_table(ws, rows, r + 1)
    for col, w in zip("ABCDEFGHIJKL", (8, 9, 5, 60, 10, 8, 8, 9, 10, 8, 12, 8)):
        ws.column_dimensions[col].width = w

    # ---------------- Interpretation ----------------
    it = wb.create_sheet("Interpretation")
    r = para(it, 1, "Interpretation with literature",
             "Each row pairs a result of this analysis with what published work in the project's reference library "
             "says. Literature statements are paraphrased from passages read in the cited sources; where the "
             "comparison is indirect (e.g. alkanes vs aromatics, acetonitrile vs water, metal vs carbon cathodes), "
             "this is noted. No claim is made beyond what the sources state.", ncol=4)
    rows = [["result here", "what the literature says", "source", "relation"],
            ["Chlorobenzenes react more slowly than bromobenzenes with similar driving forces, and the best models "
             "need a halogen-sensitive term (Wiberg C–X bond order together with a bond-energy, barrier or orbital "
             "descriptor).",
             "Alkane degradation rates on activated-carbon cathodes increased in the order chlorine < bromine < "
             "iodine (King & Mitch 2022). Higher bromide and iodide than chloride yields indicate that cleavage of "
             "C–Br and C–I bonds is favoured over C–Cl (King & Mitch 2024). In acetonitrile, bromoaromatics were "
             "chosen because their reduction waves lie inside the solvent window, unlike the analogous chlorides "
             "(Neukermans et al. 2020).",
             "KingMitch2022; KingMitch2024; Neukermans2020",
             "Consistent: a halogen-identity effect on reduction and cleavage is documented. The sources do not "
             "discuss Wiberg bond orders; its role here is inferred from the data."],
            ["Descriptors of the second reduction step (ΔG Ar• + e⁻ → Ar⁻, ΔG ArX•⁻ + e⁻ → Ar⁻ + X⁻) and the "
             "overall two-electron ΔG are among the best single descriptors and recur in multi-term models.",
             "After cleavage of the radical anion, the aryl radical of bromides and chlorides has a lower reduction "
             "potential than the parent and is reduced almost instantaneously, making the overall process a "
             "stepwise two-electron transfer (Neukermans et al. 2020). For the bromobenzenes on activated-carbon "
             "cathodes, the ΔG of the two-electron reaction was the best predictor (r² = 0.92), and multiple "
             "regressions combining radical-anion formation with fragmentation or with the second reduction were "
             "equally predictive (LaPier et al. 2026).",
             "Neukermans2020; LaPier2026",
             "Consistent, including the equivalence of the two-electron and stepwise two-descriptor models."],
            ["Radical-anion descriptors (LUMO of ArX•⁻, ΔG of radical-anion formation and fragmentation) perform "
             "well, although most computed radical anions are bent σ-type or dissociated.",
             "Aryl halides can accommodate the incoming electron in a low-energy π* orbital and therefore usually "
             "follow a stepwise mechanism; a weak C–X bond, a negative E°(RX/RX•⁻) and a positive E°(X•/X⁻) favour "
             "the concerted mechanism (Mazzucato et al. 2023). Halogenated aromatics are proposed to undergo "
             "stepwise electron transfer and bond dissociation through a radical anion, with slower C–X cleavage "
             "than alkanes (King & Mitch 2024). Chlorobenzene and bromobenzene are fast-cleaving, and both their "
             "direct and mediated reduction are kinetically controlled by the forward electron transfer "
             "(Andrieux et al. 1986, in DMF).",
             "Mazzucato2023; KingMitch2024; Andrieux1986",
             "Consistent with electron-transfer-controlled kinetics described by first-electron descriptors. The "
             "computed σ-bent or dissociated radical anions suggest fast cleavage after electron uptake, which "
             "matches 'fast-cleaving'; whether a discrete intermediate exists on the cathode is not resolved here."],
            ["Gas-phase descriptors are generally at least as predictive as SMD ones (main QSAR workbook); in the MLR "
             "screen the best-scoring set is at M06-2X/SMD but gas-phase sets are within the split-to-split spread.",
             "For halogenated alkanes the gas-phase reduction free energy gave the strongest QSAR; the authors "
             "suggest this may reflect the non-aqueous character of the activated-carbon phase, while noting that "
             "halide products must still be solvated (King & Mitch 2022). For the bromobenzenes, correlations with "
             "aqueous-phase parameters were poorer than with gas-phase ones (LaPier et al. 2026).",
             "KingMitch2022; LaPier2026",
             "Consistent with a sorbed, partly non-aqueous environment; neither level is the true environment."],
            ["Bromobenzene-trained fits overpredict the chlorobenzenes (they fall below the line) and the transfer "
             "test penalises descriptors whose range differs between halogens.",
             "A gas-phase ΔG QSAR trained on halogenated alkanes dominated by bromine-containing compounds "
             "underestimated the reactivity of fully chlorinated compounds (King & Mitch 2022). The bromobenzene QSAR "
             "was predictive for BDE-47 and BDE-99 (LaPier et al. 2026).",
             "KingMitch2022; LaPier2026",
             "Both studies show that the halogen composition of the training set matters; the direction of the "
             "Cl misprediction differs (alkanes underpredicted there, aromatics overpredicted here)."],
            ]
    r = write_table(it, rows, r + 1)
    it.cell(row=r, column=1, value="References").font = HEAD
    r += 1
    for k, v in REFS.items():
        it.cell(row=r, column=1, value=k).font = BOLD
        c = it.cell(row=r, column=2, value=v)
        c.alignment = WRAP
        it.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        r += 1
    for col, w in zip("ABCD", (45, 70, 24, 45)):
        it.column_dimensions[col].width = w

    # ---------------- Methods ----------------
    me = wb.create_sheet("Methods")
    texts = [
        ("Data", "ln k_obs of 14 compounds (7 bromobenzenes, 5 chlorobenzenes, BDE-47, BDE-99) and the computed "
                 "descriptors of the main QSAR workbook (results/correlation_data.csv). A descriptor is a candidate at a "
                 "level only if it has a value for all 14 compounds (see Candidates)."),
        ("Models", f"ln k_obs = b0 + Σ b_i·x_i with 1 to {C.MLR_MAX_TERMS} descriptors, ordinary least squares. Every "
                   "combination of candidates is fitted."),
        ("Training/test splits", f"{C.MLR_SPLITS} random splits (seed {C.MLR_SEED}), each holding out "
                                 + ", ".join(f"{k} {g}" for g, k in C.MLR_TEST_PER_GROUP.items())
                                 + " compounds as the test set and fitting on the remaining 10. Stratifying by group keeps "
                                   "every family in both sets. The model never sees its test compounds."),
        ("test RMSE", "√(mean of squared test-set errors pooled over all splits). The main ranking statistic. ± sd: "
                      "standard deviation of the per-split test RMSE; differences smaller than this are not meaningful."),
        ("Q²_F1", "1 − Σ(y − ŷ)² / Σ(y − ȳ_train)² over all test predictions: external explained variance relative to "
                  "predicting the training mean. 1 is perfect, 0 no better than the mean, < 0 worse."),
        ("train R²", "Mean R² on the training sets. A large gap between train R² and Q²_F1 signals overfitting."),
        ("R² (all 14), adj R², LOO Q²", "Fit to all compounds, the same penalised for the number of descriptors, and "
                                        "leave-one-out cross-validation on all 14."),
        ("transfer RMSE", "Fit on the bromobenzenes only, predict the chlorobenzenes and PBDEs. Tests extrapolation to "
                          "other halogens and skeletons; only defined when the bromobenzenes outnumber the parameters."),
        ("Collinearity", f"max |r|: largest correlation between two descriptors of a model; models above {C.MLR_MAX_R} "
                         "are flagged and left out of the rankings, because their coefficients are unstable. max VIF: "
                         "variance inflation factor (> 10 indicates strong collinearity)."),
        ("Standardised coefficients", "b_i·sd(x_i)/sd(y) on the All models tab: the relative weight of each descriptor."),
        ("Caveats", "With 14 compounds and up to four parameters, every model is data-limited; thousands of models "
                    "were screened, so the best score is optimistic (selection bias) even with held-out test sets. "
                    "Chlorobenzene rate constants are unpublished and have no standard errors. B3LYP/SMD lacks the "
                    "*_any descriptors of BDE-99 (radical-anion optimisation unfinished), so those are not candidates "
                    "there."),
    ]
    r = 1
    me.cell(row=r, column=1, value="Methods and statistics").font = HEAD
    r = write_table(me, [["item", "description"]] + [list(t) for t in texts], r + 1)
    me.column_dimensions["A"].width = 26
    me.column_dimensions["B"].width = 110

    # ---------------- All models ----------------
    am = wb.create_sheet("All models")
    hdr = ["level", "k", "descriptors (keys)", "descriptors", "test RMSE", "± sd", "Q²_F1", "train R²",
           "R² (all)", "adj R²", "LOO Q²", "RMSE (all)", "transfer RMSE", "max |r|", "max VIF", "collinear",
           "b0", "b1", "b2", "b3", "std b1", "std b2", "std b3"]
    rows = [hdr]
    for m in models:
        cf, sc = m["coef"] + [None] * 4, m["std_coef"] + [None] * 3
        rows.append([QW.SHORT[m["level"]], m["k"], "; ".join(m["descriptors"]), model_name(m), m["RMSE_test"],
                     m["RMSE_test_sd"], m["Q2_F1"], m["R2_train"], m["R2"], m["adj_R2"], m["Q2_LOO"], m["RMSE"],
                     m["RMSE_transfer"], m["max_abs_r"], min(m["max_VIF"], 1e6), "yes" if m["collinear"] else "no",
                     *cf[:4], *sc[:3]])
    for i, row in enumerate(rows):
        am.append(row)
    for c in am[1]:
        c.font, c.fill = BOLD, SHADE
    for row in am.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.000"
    am.freeze_panes = "E2"
    am.auto_filter.ref = am.dimensions
    for col, w in zip(range(1, 24), (7, 4, 40, 55) + (9,) * 19):
        am.column_dimensions[get_column_letter(col)].width = w

    # ---------------- Candidates ----------------
    ca = wb.create_sheet("Candidates")
    rows = [["level", "candidate descriptors (values for all 14 compounds)", "excluded (missing for some compounds)"]]
    for lv, i in info.items():
        rows.append([QW.SHORT[lv], "; ".join(label(d) for d in i["candidates"]), "; ".join(label(d) for d in i["excluded"])])
    write_table(ca, rows, 1)
    for col, w in zip("ABC", (8, 90, 90)):
        ca.column_dimensions[col].width = w

    # ---------------- Parity charts ----------------
    jobs = {}
    picks = [("Parity best", best), ("Parity best 2-desc", best_k[2]), ("Parity best single", best_k[1]),
             ("Parity best transfer", transfer[0])]
    seen = set()
    for name, m in picks:
        key = (m["level"], tuple(m["descriptors"]))
        if key in seen:
            continue
        seen.add(key)
        parity_sheet(wb, name, data, m, st, jobs, preview_dir)

    for w_ in (ws, it, me, ca):
        QW.fit_heights(w_)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    X.patch_many(out, st, jobs)
    print(f"wrote {out}: {len(models)} models; best {model_name(best)} ({best['level']}) "
          f"test RMSE {best['RMSE_test']:.2f}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "figures/QSAR_MLR_exploration.xlsx",
         sys.argv[2] if len(sys.argv) > 2 else None)
